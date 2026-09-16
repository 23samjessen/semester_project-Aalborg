from __future__ import annotations

from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from make_member_3_4_guide import (
    FIG_DIR,
    REPORT_DIR,
    Guide,
    add_page_field,
    replace_placeholder_with_links,
)


def set_teacher_footer(guide: Guide) -> None:
    footer = guide.doc.sections[0].footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Public Blocklist Study – Project Plan  |  ")
    run.font.name = "Liberation Serif"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(90, 90, 90)
    add_page_field(footer)


def add_cover(guide: Guide) -> None:
    doc = guide.doc
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(75)
    r = p.add_run("Project Plan for the Supervisor")
    r.bold = True
    r.font.name = "Liberation Serif"
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(31, 78, 121)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("Analysis of Public Blocklists from a Danish Perspective")
    r.font.name = "Liberation Serif"
    r.font.size = Pt(15)
    r.bold = True

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(25)
    r = p3.add_run("Scope, method, work division, schedule and planned deliverables")
    r.font.name = "Liberation Serif"
    r.font.size = Pt(11.5)
    r.font.italic = True

    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(80)
    for line in [
        "Group members: [Member 1] · [Member 2] · [Member 3] · [Member 4] · [Member 5] · [Member 6]",
        "Supervisor: [Name]",
        "Aalborg University – Cyber Security",
        "Version: 16 September 2026",
    ]:
        r = p4.add_run(line + "\n")
        r.font.name = "Liberation Serif"
        r.font.size = Pt(10.5)

    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p5.paragraph_format.space_before = Pt(80)
    r = p5.add_run("This plan is a proposal for discussion and approval. Dates and source choices can be updated with the supervisor.")
    r.font.name = "Liberation Serif"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(90, 90, 90)


def build_plan() -> Path:
    guide = Guide()
    set_teacher_footer(guide)
    add_cover(guide)

    guide.heading("Table of Contents", 1, "toc")
    guide.paragraph("The entries below are linked to the corresponding sections in the plan.")
    guide.doc.add_paragraph("[[TOC]]")

    guide.heading("1. Executive summary", 1, "executive_summary")
    guide.paragraph(
        "The group proposes a small empirical study of public cybersecurity blocklists from a Danish perspective. We will examine how selected providers describe their listing and removal processes, collect dated snapshots of their public feeds, identify observations with defensible Danish-related evidence, and measure changes over time. The study combines a theoretical literature review with a reproducible hands-on data pipeline."
    )
    guide.paragraph(
        "The practical output will be a documented dataset, source and policy register, normalized observations, Danish classification evidence, descriptive tables, visualizations and a final report. The project will not scan public hosts, generate malicious traffic, submit false abuse reports or claim that list membership alone proves current maliciousness."
    )
    guide.table(
        ["Item", "Planned decision"],
        [
            ("Study design", "Longitudinal observational study plus provider-document analysis"),
            ("Minimum collection period", "At least 30 distinct calendar dates, normally one snapshot per day"),
            ("Initial sources", "FireHOL Level 2, Spamhaus DROP IPv4 and blocklist.de all.txt"),
            ("Optional source", "URLhaus, only after the group approves a separate domain/URL scope and checks access terms"),
            ("Primary unit of analysis", "Indicator–source–time observation"),
            ("Main output", "Evidence-based comparison of listing dynamics, overlap, persistence and correction procedures"),
        ],
        widths=[1.85, 4.8],
        font_size=9.2,
    )

    guide.heading("2. Background and motivation", 1, "background")
    guide.paragraph(
        "Blocklists are used by network defenders to reject traffic, generate alerts or increase risk scores. They are not all designed for the same purpose: some publish aggregated network ranges, some publish curated high-risk ranges and some represent reports of observed attacks. Their contents can change when evidence is updated, a network is reassigned, a domain expires or an operator changes a policy."
    )
    guide.paragraph(
        "This creates a practical and scientific problem. A list may be useful for a defender while still being difficult to interpret for an address owner. Without timestamps, source semantics and uncertainty labels, it is easy to confuse historical listing with current maliciousness, country allocation with physical location, or a broad CIDR range with a single responsible host."
    )
    guide.paragraph(
        "The Danish perspective gives the group a concrete attribution question: which listed indicators have evidence associated with Denmark, how stable is that evidence, and how should conflicting or incomplete evidence be reported?"
    )

    guide.heading("3. Problem formulation, aim and research questions", 1, "problem_and_questions")
    guide.heading("3.1 Problem formulation", 2, "problem_formulation")
    guide.paragraph(
        "Public blocklists differ in purpose, transparency, representation and update behavior. It is therefore unclear how consistently their published contents identify, retain and remove indicators associated with Denmark, and how difficult it is for a benign owner to understand or correct a listing."
    )
    guide.heading("3.2 Aim", 2, "aim")
    guide.paragraph(
        "The aim is to characterize selected public blocklists as changing security information systems and to evaluate, with a reproducible dataset and documented source analysis, what their public data can and cannot support about Danish-related indicators, listing dynamics and correction procedures."
    )
    guide.heading("3.3 Research questions", 2, "research_questions")
    guide.table(
        ["RQ", "Question", "Planned evidence"],
        [
            ("RQ1", "How do the selected blocklists describe and implement listing, updating, expiry and removal?", "Provider documentation, feed fields, update behavior and policy pages."),
            ("RQ2", "What patterns of listing, persistence, turnover and cross-list overlap are observed for Danish-related indicators during the collection period?", "At least 30 days of timestamped snapshots, Danish evidence and descriptive analysis."),
            ("RQ3", "What practical and procedural barriers affect correction or delisting of a benign owned asset?", "Published correction procedures and, only if permitted, a controlled owned-asset case."),
        ],
        widths=[0.55, 3.1, 3.0],
        font_size=8.9,
    )

    guide.heading("4. Scope and delimitations", 1, "scope")
    guide.table(
        ["Included", "Delimitation"],
        [
            ("Public feed snapshots", "Only the selected public routes and the dates successfully collected are analyzed."),
            ("IP addresses and CIDR ranges", "They are retained as separate representations; CIDRs are not silently expanded into individual IPs."),
            ("Domains and URLs", "If URLhaus is enabled, domains/URLs will be analyzed separately from IP/CIDR indicators."),
            ("Danish-related evidence", "Allocation, geolocation, ASN, DNS and other signals are recorded separately with uncertainty."),
            ("Provider policies", "The group will analyze publicly documented procedures and source semantics."),
            ("Active security testing", "No public-host scanning, exploitation, false reporting or malicious traffic generation."),
            ("Local Wi-Fi address", "192.168.0.161 is not a project target and is not used by the collection program."),
        ],
        widths=[2.0, 4.65],
        font_size=9.0,
    )
    guide.paragraph("The study will not generalize from these feeds to the whole Internet or treat a provider’s listing as proof of present-day maliciousness. These limitations will be stated in the report’s methodology and discussion.")

    guide.heading("5. Methodology", 1, "methodology")
    guide.paragraph("The project uses a transparent, longitudinal workflow. Each transformation must be traceable to an original timestamped source snapshot and each interpretation must state the evidence and uncertainty supporting it.")
    pipeline = FIG_DIR / "figure_1_research_pipeline.png"
    if pipeline.exists():
        guide.figure(pipeline, "Figure 1", "Planned research and data pipeline from source retrieval to final interpretation.", width=6.35)
    guide.heading("5.1 Work package A — literature and theory", 2, "method_literature")
    guide.paragraph("Member 2 will establish the theoretical framework through a documented literature search. The review will cover open blocklist transparency, list dynamics, persistence, false-positive risk, country attribution and responsible use. Each source will be summarized, assessed for relevance and cited in IEEE style. The group will use the literature to define concepts and explain the research gap, not merely to provide background paragraphs.")
    guide.heading("5.2 Work package B — source dissection and policy analysis", 2, "method_source_dissection")
    guide.paragraph("The group will inspect each selected source’s purpose, representation, fields, update frequency, retention/expiry behavior, access route, terms of use and correction or delisting process. We will record the exact URL, access date, provider wording and any uncertainty. Provider semantics will be kept separate: FireHOL, Spamhaus DROP and blocklist.de are not interchangeable labels for the same kind of evidence.")
    guide.heading("5.3 Work package C — data collection and provenance", 2, "method_collection")
    guide.paragraph("Member 3 will run the complete daily pipeline once per day at an agreed UTC time. The shell collector downloads the public feed files using curl, stores immutable raw snapshots under data/raw/YYYY-MM-DD/, records HTTP status and byte count, and calculates SHA-256 hashes. The runner creates a unique run_records/<run-id>/ folder containing the UTC start/end times, commands, exit codes, per-run collection log, parser notes, command output and a copy of the analysis. The current collector uses public HTTPS feed URLs rather than authenticated APIs. URLhaus remains disabled unless separately approved.")
    guide.heading("5.4 Work package D — normalization and Danish evidence", 2, "method_normalization")
    guide.paragraph("The normalizer converts the different file formats into one CSV while preserving source path, source value, timestamp and hash. IP, CIDR, domain and URL indicators remain distinct. Member 4 will add passive Danish evidence from permitted sources, record source and retrieval time, and apply the agreed labels DK-supported, DK-weak, Conflicting or Unknown.")
    guide.heading("5.5 Work package E — quantitative analysis", 2, "method_analysis")
    guide.paragraph("Member 5 will analyze the normalized observations. Planned measures are unique indicators per source/day, additions and removals between successful snapshots, persistence across snapshot days, exact representation-level overlap and Jaccard similarity. Results will be stratified by source and indicator type. The daily runner also calls a separate comparison module that writes two-date summaries and individual change files under comparisons/<run-id>/. Results will not be overwritten by the comparison step. The analysis is descriptive; causal or operational claims require discussion of source semantics and limitations.")
    guide.heading("5.6 Work package F — interpretation and conclusion", 2, "method_interpretation")
    guide.paragraph("Member 6 will lead the policy, ethics, discussion and conclusion work with input from all members. The final conclusion will answer RQ1–RQ3 using only verified tables, figures, documented policies and recorded evidence. Unresolved cases will be reported as uncertainty rather than forced into a definitive classification.")

    guide.heading("6. Exact steps we will follow", 1, "step_by_step")
    guide.table(
        ["Step", "What the group will do", "Evidence/output", "Lead"],
        [
            ("1. Approve the design", "Confirm RQs, source set, dates, UTC time, definitions, Danish rule and ethics boundaries.", "Signed-off method sheet and contribution plan.", "Member 1 + all"),
            ("2. Build the theory", "Search, read and compare literature; define key terms and research gap.", "Literature matrix and IEEE reference register.", "Member 2"),
            ("3. Dissect providers", "Read feed documentation and correction policies; inspect actual formats.", "Source/policy matrix and cited captures.", "Member 6 + Member 2"),
            ("4. Test the pipeline", "Create the Python environment and run a baseline collection, parser and analysis.", "Successful smoke-test output and issue log.", "Member 3"),
            ("5. Collect daily", "Run the complete pipeline once per day and preserve raw snapshots, hashes, timestamps and command outcomes.", "At least 30 distinct dates and one run record per execution.", "Member 3"),
            ("6. Normalize", "Convert raw formats into one provenance-preserving observation table.", "observations.csv and parser notes.", "Member 3"),
            ("7. Classify Denmark", "Retrieve permitted passive evidence and label candidates with uncertainty.", "Danish evidence CSV and classification CSV.", "Member 4"),
            ("8. Analyze", "Calculate size, additions, removals, persistence and overlap; inspect missing days.", "CSV result tables and PNG figures.", "Member 5"),
            ("9. Compare", "Compare the latest snapshots from two distinct days using exact canonical indicators.", "Separate comparison manifest, summaries, change files and figure.", "Members 3 + 5"),
            ("10. Discuss", "Relate observations to literature, policies, ethics and validity threats.", "Discussion, limitations and recommendations.", "Member 6 + all"),
            ("11. Conclude and defend", "Answer the research questions, verify references and prepare the oral presentation.", "Final report, appendices, contribution log and presentation.", "Member 1 + all"),
        ],
        widths=[0.9, 2.55, 2.25, 0.95],
        font_size=8.1,
    )

    guide.heading("7. Six-member division of work", 1, "work_division")
    guide.table(
        ["Member", "Primary responsibility", "Concrete deliverables", "Review responsibility"],
        [
            ("1", "Coordination and report integration", "Problem statement, aim, RQs, meeting log, final structure and oral defence plan.", "Check coherence and that every chapter answers the RQs."),
            ("2", "Literature review and IEEE sources", "Search log, literature matrix, theoretical framework and reference register.", "Check source quality and citation consistency."),
            ("3", "Collection and normalization", "Raw snapshots, collection log, hashes, parser, normalized CSV and reproducibility notes.", "Check data integrity and repeatability."),
            ("4", "Danish classification and enrichment", "Evidence table, passive enrichment, classification labels and uncertainty notes.", "Check attribution, privacy and overclaiming."),
            ("5", "Quantitative analysis and visualization", "Daily summary, persistence, turnover, overlap analysis, tables and figures.", "Check that results match the method and data."),
            ("6", "Policy, ethics, discussion and conclusion", "Provider policy comparison, delisting analysis, limitations, recommendations and conclusion.", "Check responsible interpretation and validity threats."),
        ],
        widths=[0.55, 1.65, 2.5, 1.95],
        font_size=8.15,
    )
    guide.paragraph("The work packages are assigned for ownership, not isolation. Each member will review the full report and understand how the raw snapshots become the final findings. Members 3 and 4 jointly own the main hands-on evidence work, with Member 5 using their verified outputs for the quantitative chapter.")

    guide.heading("8. Schedule and milestones", 1, "schedule")
    guide.paragraph("The collection period is planned as 30 distinct dates. The academic writing and source work run in parallel so that the final report is not delayed until the last collection day.")
    timeline = FIG_DIR / "figure_4_collection_timeline.png"
    if timeline.exists():
        guide.figure(timeline, "Figure 2", "Proposed 30-day collection timeline and review checkpoints.", width=6.35)
    guide.table(
        ["Period", "Main activities", "Milestone for supervisor"],
        [
            ("Week 1", "Approve scope and RQs; assign roles; start literature and policy register; test the collector and parser.", "Method and ethics boundaries agreed."),
            ("Week 2", "Start daily snapshots; verify source status and hashes; draft introduction, context and methodology.", "Baseline and first quality review."),
            ("Week 3", "Continue collection; inspect format stability; build candidate list for Danish evidence; continue literature synthesis.", "No silent failures; candidate-selection rule approved."),
            ("Week 4", "Midpoint analysis; begin evidence classification; draft preliminary tables and figures; review policy procedures.", "Midpoint data and analysis meeting."),
            ("Week 5", "Continue collection; resolve conflicting/unknown cases; draft findings and discussion structure.", "Policy matrix and uncertainty review."),
            ("Week 6", "Complete 30-day coverage; freeze raw data; rerun normalization and analysis; complete figures and appendices.", "Data freeze and final findings review."),
            ("Final phase", "Integrate report, verify IEEE citations, cross-review all claims, write conclusion and prepare defence.", "Submission-ready report and presentation."),
        ],
        widths=[1.0, 3.75, 1.9],
        font_size=8.5,
    )

    guide.heading("9. Current status and next actions", 1, "current_status")
    guide.paragraph("The first baseline run has been completed. FireHOL, Spamhaus DROP and blocklist.de returned successfully. The normalizer reported 89,277 observations and 0 parser notes. The analysis reported 3 daily source rows, 44,746 persistence rows and 2 overlap rows. URLhaus is intentionally disabled. These results demonstrate that the pipeline runs; they are not yet the final historical findings because only one calendar date is represented.")
    guide.table(
        ["Next action", "Responsible", "Success condition"],
        [
            ("Agree on one UTC collection time and start/end dates.", "Member 1 + all", "The schedule is written in the method log."),
            ("Continue the daily three-feed collection.", "Member 3", "Every day has status, paths and hashes or a documented failure."),
            ("Review normalized schema and candidate indicators.", "Members 3 and 4", "Member 4 receives a clean, traceable candidate list."),
            ("Begin Danish evidence capture on a documented sample.", "Member 4", "Every classification has source, timestamp and uncertainty."),
            ("Confirm the analysis plan and figure list.", "Member 5", "Each planned metric maps to an RQ and a result file."),
            ("Approve provider-policy and safe-delisting scope.", "Member 6 + supervisor", "No active testing occurs without explicit authorization."),
        ],
        widths=[3.1, 1.35, 2.2],
        font_size=8.7,
    )

    guide.heading("10. Data management, quality and reproducibility", 1, "quality")
    guide.bullet("Raw snapshots are immutable evidence. They will not be edited or overwritten after download.")
    guide.bullet("The collection log records retrieval time, source, URL, HTTP status, bytes, SHA-256, path, status and notes.")
    guide.bullet("Derived CSV files can be regenerated from raw snapshots using the project scripts.")
    guide.bullet("Each complete run has a separate run manifest, event timeline, command logs and archived analysis copy.")
    guide.bullet("Two-date comparisons are stored under comparisons/ and never overwrite results/.")
    guide.bullet("If a source snapshot is missing on one selected date, its comparison counts are marked NA rather than reported as removals.")
    guide.bullet("Missing days, failed downloads, duplicate same-day runs and provider format changes will be logged explicitly.")
    guide.bullet("All country and enrichment evidence will include source, access/retrieval date and uncertainty.")
    guide.bullet("The group will separate indicator type, source semantics and time rather than combining unlike representations.")
    guide.bullet("Every numerical claim in the report will point to a result table/figure and its underlying evidence.")
    guide.bullet("No API keys, private correspondence or personal data will be stored in the repository or report.")

    guide.heading("11. Ethics, safety and legal boundaries", 1, "ethics")
    guide.paragraph("The project is based on passive observation of publicly published feed data and provider documentation. Public availability does not automatically grant permission to scan, connect to or test a listed host. The group will not generate malicious traffic, perform exploitation, submit false abuse reports or attempt to manipulate a provider’s list.")
    guide.paragraph("If a delisting or correction case is included, it must involve an asset owned or explicitly controlled by the group, follow the provider’s documented process, avoid deception and be approved by the supervisor before any action. Otherwise, the group will answer RQ3 through policy analysis and documented public procedures.")
    guide.paragraph("The Danish classification will be presented as evidence-supported and uncertain where appropriate. Allocation country, geolocation and reverse DNS are different signals; none alone proves ownership, physical location or malicious intent.")

    guide.heading("12. Expected deliverables and acceptance criteria", 1, "deliverables")
    guide.table(
        ["Deliverable", "Acceptance criterion"],
        [
            ("Project report", "Contains problem, theory, literature review, methodology, findings, discussion, conclusion, IEEE references, index and appendices."),
            ("Literature review", "Uses relevant sources, explains the research gap and cites every borrowed idea in IEEE style."),
            ("Provider dissection", "Documents purpose, format, update behavior, semantics and correction/delisting procedures for each selected source."),
            ("Raw dataset", "Contains the agreed collection window, source paths, timestamps, status and hashes."),
            ("Run audit package", "Contains one timestamped run record per execution with commands, outcomes, parser notes and archived analysis."),
            ("Comparison package", "Contains independent previous/current-day manifests, source summaries, individual changes and figures."),
            ("Normalized dataset", "Preserves provenance and distinguishes IP, CIDR, domain and URL representations."),
            ("Danish evidence", "Includes evidence source, retrieval time, classification rule and uncertainty/conflict notes."),
            ("Analysis package", "Includes daily summary, additions/removals, persistence, overlap/Jaccard tables and readable figures."),
            ("Contribution record", "Shows the work performed by all six members and the cross-review process."),
            ("Presentation/defence", "Every member can explain the problem, method, own contribution, limitations and main evidence."),
        ],
        widths=[1.85, 4.8],
        font_size=8.8,
    )

    guide.heading("13. Proposed final report structure", 1, "report_structure")
    guide.numbered("Introduction: context, problem formulation, aim, research questions and contribution.")
    guide.numbered("Technical background: blocklists, feeds, DNSBLs, IP/CIDR, domains/URLs, country attribution and provider semantics.")
    guide.numbered("Literature review: prior work on transparency, dynamics, effectiveness, false positives and correction.")
    guide.numbered("Methodology: source selection, collection protocol, data model, Danish rule, analysis measures, ethics and validity.")
    guide.numbered("Findings: verified collection coverage, source characteristics, Danish classifications, trends, persistence, turnover and overlap.")
    guide.numbered("Discussion: interpretation against the literature, practical significance, policy implications and threats to validity.")
    guide.numbered("Conclusion: direct answers to RQ1–RQ3, limitations and future work.")
    guide.numbered("References: IEEE numbered bibliography and source access information.")
    guide.numbered("Appendices: commands, data dictionary, source/policy matrix, collection log summary, figures and contribution log.")

    guide.heading("14. Questions for supervisor approval", 1, "supervisor_questions")
    for item in [
        "Is the three-source baseline sufficient, or should URLhaus be added as a separate domain/URL source?",
        "Is a 30-distinct-day collection period appropriate for the module schedule?",
        "Does the proposed DK-supported/DK-weak/Conflicting/Unknown evidence rule need adjustment?",
        "Which passive enrichment services and rate limits are acceptable for a small candidate sample?",
        "Should the group include an owned-asset correction case, or use policy analysis only?",
        "Which figures and statistical summaries are sufficient for the expected academic level?",
        "Are the proposed chapter ownership and six-member contribution records appropriate?",
    ]:
        guide.bullet(item)
    guide.paragraph("The group will update this plan after the supervisor’s feedback and record the agreed changes in the method log.")

    guide.heading("15. Preliminary IEEE-style references and official sources", 1, "references")
    references = [
        "[1] A. Feal et al., “Blocklist Babel: On the Transparency and Dynamics of Open Source Blocklisting,” IEEE Trans. Netw. Serv. Manag., vol. 18, no. 2, pp. 1334–1349, Jun. 2021, doi: 10.1109/TNSM.2021.3075552.",
        "[2] L. Deri and F. Fusco, “Evaluating IP Blacklists Effectiveness,” arXiv preprint arXiv:2308.08356, 2023, doi: 10.48550/ARXIV.2308.08356.",
        "[3] FireHOL, “IP Lists and Blocklist IP Sets,” [Online]. Available: https://iplists.firehol.org/ and https://github.com/firehol/blocklist-ipsets/.",
        "[4] Spamhaus, “Do Not Route Or Peer (DROP),” [Online]. Available: https://www.spamhaus.org/blocklists/do-not-route-or-peer/.",
        "[5] blocklist.de, “Export Lists,” [Online]. Available: https://www.blocklist.de/en/export.html.",
        "[6] URLhaus, “API and Feeds,” [Online]. Available: https://urlhaus.abuse.ch/api/ and https://urlhaus.abuse.ch/feeds/.",
        "[7] Aalborg University, “Distributed Systems Security module,” [Online]. Available: https://moduler.aau.dk/course/2026-2027/ESNCYSK1P1.",
    ]
    for reference in references:
        guide.paragraph(reference)
    guide.paragraph("The final report will update this preliminary list with the sources actually used and the access dates required by the chosen citation format.")

    guide.heading("16. Short message to send with this plan", 1, "teacher_message")
    guide.paragraph(
        "Dear [Supervisor], our group proposes to study selected public blocklists from a Danish perspective. We will combine a literature review and provider-policy analysis with a 30-day reproducible collection of FireHOL, Spamhaus DROP and blocklist.de snapshots. We will preserve raw files and hashes, normalize the different feed formats, document Danish-related evidence with uncertainty, and analyze listing size, additions, removals, persistence and overlap. Each daily execution will have a timestamped run record, and each two-date comparison will be stored separately from the cumulative results. We will not scan public hosts or treat list membership as proof of current maliciousness. The attached plan describes the work division, milestones, deliverables and the points on which we would appreciate your approval."
    )

    toc_entries = [(title, level, anchor) for title, level, anchor in guide.headings if title != "Table of Contents"]
    replace_placeholder_with_links(guide.doc, "[[TOC]]", toc_entries)

    output = REPORT_DIR / "Blocklist_Project_Plan_for_Supervisor.docx"
    guide.doc.save(output)
    print(output)
    return output


if __name__ == "__main__":
    build_plan()
