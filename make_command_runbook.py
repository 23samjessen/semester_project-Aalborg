from __future__ import annotations

import csv
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from make_member_3_4_guide import (
    FIG_DIR,
    REPORT_DIR,
    SCREENSHOT,
    Guide,
    add_code,
    add_page_field,
    replace_placeholder_with_links,
)


ROOT = Path(__file__).resolve().parent


def set_runbook_footer(guide: Guide) -> None:
    footer = guide.doc.sections[0].footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Blocklist Project Command Runbook  |  ")
    run.font.name = "Liberation Serif"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(90, 90, 90)
    add_page_field(footer)


def add_cover(guide: Guide) -> None:
    doc = guide.doc
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    r = p.add_run("Blocklist Project Command Runbook")
    r.bold = True
    r.font.name = "Liberation Serif"
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(31, 78, 121)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("Analysis of Public Blocklists from a Danish Perspective")
    r.font.name = "Liberation Serif"
    r.font.size = Pt(15)
    r.bold = True

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(25)
    r = p3.add_run("Copy-paste commands, expected outputs and the daily research workflow")
    r.font.name = "Liberation Serif"
    r.font.size = Pt(11.5)
    r.font.italic = True

    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(90)
    for line in [
        "Prepared for the six-person project group",
        "Aalborg University – Cyber Security",
        "16 September 2026",
    ]:
        r = p4.add_run(line + "\n")
        r.font.name = "Liberation Serif"
        r.font.size = Pt(11)

    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p5.paragraph_format.space_before = Pt(85)
    r = p5.add_run(
        "The commands operate on public feed files and local project data. They do not scan your Mac or public hosts."
    )
    r.font.name = "Liberation Serif"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(90, 90, 90)


def current_counts() -> dict[str, object]:
    counts: dict[str, object] = {
        "normalized": 0,
        "daily": 0,
        "persistence": 0,
        "overlap": 0,
        "parser_notes": 0,
        "daily_rows": [],
    }

    normalized = ROOT / "data" / "normalized" / "observations.csv"
    if normalized.exists():
        with normalized.open(newline="", encoding="utf-8") as handle:
            counts["normalized"] = sum(1 for _ in csv.DictReader(handle))

    errors = ROOT / "logs" / "normalization_errors.csv"
    if errors.exists():
        with errors.open(newline="", encoding="utf-8") as handle:
            counts["parser_notes"] = sum(1 for _ in csv.DictReader(handle))

    for key, filename in (("daily", "daily_summary.csv"), ("persistence", "persistence.csv"), ("overlap", "overlap.csv")):
        path = ROOT / "results" / filename
        if path.exists():
            with path.open(newline="", encoding="utf-8") as handle:
                rows = list(csv.DictReader(handle))
            counts[key] = len(rows)
            if key == "daily":
                counts["daily_rows"] = rows
    return counts


def build_runbook() -> Path:
    guide = Guide()
    set_runbook_footer(guide)
    add_cover(guide)

    guide.heading("Table of Contents", 1, "toc")
    guide.paragraph("Every entry is linked. Select a command section to jump to it.")
    guide.doc.add_paragraph("[[TOC]]")

    guide.heading("1. Start here", 1, "start_here")
    guide.paragraph(
        "This file is the operational companion to the project guide. It tells you exactly what to type in Terminal, what each command changes, what output to expect and what to report to the group. Run commands from the project directory, the folder that contains scripts/, data/, results/, run_records/, comparisons/ and requirements.txt."
    )
    guide.heading("1.1 Confirm the project directory", 2, "confirm_directory")
    guide.paragraph("If your terminal prompt already ends in Blocklist_Danish_Project_Package, you are probably in the correct directory. Confirm it before running anything:")
    add_code(guide.doc, r"""pwd
ls
find . -maxdepth 2 -type f -print | sort""")
    guide.paragraph(
        "You should see requirements.txt, scripts/collect_snapshots.sh, scripts/normalize_snapshot.py, scripts/analyze_blocklists.py, scripts/run_daily_pipeline.py, data/, results/, run_records/, comparisons/ and logs/. If you are elsewhere, change the path in the next command to the real location of the downloaded project."
    )
    add_code(guide.doc, r"""cd /absolute/path/to/Blocklist_Danish_Project_Package
pwd""")
    guide.paragraph(
        "Your Wi-Fi address, 192.168.0.161, is a local address on your Mac. Nothing in this workflow uses it as a target. The collector downloads public feed files; it is not a network scanner."
    )

    guide.heading("2. Create and activate the Python environment", 1, "environment_setup")
    guide.paragraph("These are the setup commands you ran. They create an isolated environment and install the libraries required by the project.")
    add_code(guide.doc, r"""python3 -m venv .venv
source .venv/bin/activate
python -m pip install -r requirements.txt""")
    guide.paragraph("The prompt should begin with (.venv) after activation. The warning about a damaged or ignored pip cache is not a project failure when installation finishes with Successfully installed.")
    guide.heading("2.1 Recommended verification", 2, "verify_environment")
    add_code(guide.doc, r"""which python
python --version
python -m pip --version
python -m pip check""")
    guide.paragraph("The first path should point inside the project’s .venv directory. pip check should report that the installed packages have compatible dependencies. When you finish working, leave the environment with:")
    add_code(guide.doc, "deactivate")
    guide.paragraph("If you open a new Terminal window later, run source .venv/bin/activate again before using the Python scripts.")

    guide.heading("3. Collect public blocklist snapshots", 1, "collect_snapshots")
    guide.paragraph("Collection is Member 3’s main operation. The shell script uses curl to make ordinary HTTPS downloads, saves each original response under data/raw/YYYY-MM-DD/, calculates a SHA-256 hash and appends status information to logs/collection_log.csv.")
    guide.heading("3.1 Recommended complete daily command", 2, "daily_pipeline_command")
    add_code(guide.doc, r"""export BLOCKLIST_USER_AGENT='AAU-blocklist-study/0.1 contact: YOUR-GROUP-EMAIL'
python scripts/run_daily_pipeline.py""")
    guide.paragraph("Use this command once per collection day. It runs collection, normalization, cumulative analysis and day-to-day comparison in order. It creates a new run_records/<run-id>/ folder before starting, so the command output and timestamps can be reviewed even if a later step fails.")
    guide.heading("3.2 Manual collection command", 2, "collect_command")
    add_code(guide.doc, r"""export BLOCKLIST_USER_AGENT='AAU-blocklist-study/0.1 contact: YOUR-GROUP-EMAIL'
bash scripts/collect_snapshots.sh""")
    guide.paragraph("Replace YOUR-GROUP-EMAIL with a monitored group address. A useful user agent identifies the study and gives providers a way to contact the group. Do not use a private password, API key or personal secret in this value.")
    guide.heading("3.3 What the current collector connects to", 2, "collector_sources")
    guide.table(
        ["Source", "Route used by the script", "Meaning"],
        [
            ("FireHOL", "Public raw GitHub feed: firehol_level2.netset", "Aggregated IP/CIDR indicators."),
            ("Spamhaus DROP", "Official drop_v4.json feed", "Curated IPv4 network ranges and provider metadata."),
            ("blocklist.de", "Official lists/blocklist.de/lists/all.txt export", "Reported-attack IP indicators."),
            ("URLhaus", "Optional; disabled in the current run", "A separate domain/URL malware dimension."),
        ],
        widths=[1.25, 2.8, 2.45],
        font_size=9.0,
    )
    guide.paragraph("These are direct public feed URLs, not authenticated API calls. The current successful run uses no API key. The Python package requests is installed for possible future enrichment, but the collector itself calls curl from the shell script. AbuseIPDB, GreyNoise, VirusTotal and RIPEstat are not called by the current collector.")
    guide.heading("3.4 Expected current output", 2, "collect_expected_output")
    add_code(guide.doc, "firehol: ok (... bytes)\nspamhaus: ok (... bytes)\nblocklist_de: ok (... bytes)\nURLhaus is disabled. Set FETCH_URLHAUS=1 and URLHAUS_URL after checking its current access route and terms.")
    guide.paragraph("The byte count can change each day. What matters is that the status is ok, a file path is recorded and a non-zero hash is stored. URLhaus being disabled is a configuration message, not a failure of the three-source baseline.")
    guide.heading("3.5 Optional URLhaus collection", 2, "urlhaus_optional")
    guide.paragraph("Do not enable this just to make the output larger. First ask the group to approve the URLhaus scope, check the current official access route and terms, and decide whether domains/URLs will be analyzed separately from IPs/CIDRs.")
    add_code(guide.doc, r"""export FETCH_URLHAUS=1
export URLHAUS_URL='https://urlhaus.abuse.ch/downloads/csv_recent/'
bash scripts/collect_snapshots.sh""")
    guide.paragraph("If an authenticated route is required, configure the key only through the shell environment expected by the script. Never commit a key to the project or paste it into the report.")
    guide.heading("3.6 Check the collection log", 2, "collection_log_check")
    add_code(guide.doc, r"""column -s, -t < logs/collection_log.csv | tail -n 10
find data/raw -type f -print | sort""")
    guide.paragraph("Review source, URL, HTTP status, byte count, SHA-256, path, status and notes. Keep every raw snapshot, including a failed or duplicate run, until the group has documented what happened. Do not edit raw files by hand.")

    guide.heading("4. Normalize the raw snapshots", 1, "normalize_snapshots")
    guide.paragraph("Normalization is Member 3’s second main operation. The Python script reads the different provider formats and writes one common CSV. It preserves the original path, source value, retrieval time and file hash so that every row can be traced back to a raw snapshot.")
    guide.heading("4.1 The command to run", 2, "normalize_command")
    add_code(guide.doc, r"""python scripts/normalize_snapshot.py \
  --raw-dir data/raw \
  --out data/normalized/observations.csv \
  --errors logs/normalization_errors.csv""")
    guide.heading("4.2 Check the parser result", 2, "normalize_checks")
    add_code(guide.doc, r"""cat logs/normalization_errors.csv
head -n 3 data/normalized/observations.csv
wc -l data/normalized/observations.csv
cut -d, -f3 data/normalized/observations.csv | cut -c1-10 | sort -u""")
    guide.paragraph("A file containing only the header path,error means the parser recorded zero notes. The row count from wc -l includes the CSV header, while the script’s wrote N normalized observations message counts data rows. The final command lists the calendar dates currently represented in the normalized file; it must eventually show at least 30 distinct dates.")
    guide.paragraph("The normalizer keeps IP addresses, CIDR ranges, domains and URLs as different indicator types. It canonicalizes equivalent IP/CIDR text, lower-cases domain names and keeps URLs as URLs. It does not decide whether an indicator is Danish or malicious.")

    guide.heading("5. Analyze the normalized observations", 1, "analyze_observations")
    guide.paragraph("Analysis is Member 5’s main operation, but Member 3 should run it after every collection so the group can see whether the pipeline is healthy. The script creates descriptive CSV tables and PNG figures; it does not produce a final causal claim automatically.")
    guide.heading("5.1 The command to run", 2, "analyze_command")
    add_code(guide.doc, r"""python scripts/analyze_blocklists.py \
  --input data/normalized/observations.csv \
  --out-dir results""")
    guide.heading("5.2 Inspect the generated results", 2, "analyze_checks")
    add_code(guide.doc, r"""column -s, -t < results/daily_summary.csv
head -n 5 results/persistence.csv
cat results/overlap.csv
cat results/analysis_notes.txt
find results -maxdepth 1 -type f -print | sort""")
    guide.table(
        ["Output", "What it means", "How it is used"],
        [
            ("daily_summary.csv", "Unique indicators, additions, removals and raw rows per source/day.", "Trend tables and time-series figures."),
            ("persistence.csv", "First/last observation and the fraction of successful snapshot days present.", "Persistence analysis after multiple dates."),
            ("overlap.csv", "Exact canonical intersections, unions and Jaccard values by same indicator type.", "Source comparison; IP/CIDR types stay separate."),
            ("daily_unique_indicators.png", "Unique canonical indicators over time.", "Results figure."),
            ("daily_additions_removals.png", "New and absent indicators relative to the previous successful snapshot.", "Turnover figure."),
            ("source_overlap.png", "Overlap visualized for comparable representations.", "Comparison figure."),
        ],
        widths=[1.75, 2.95, 1.85],
        font_size=8.8,
    )
    guide.paragraph("If Matplotlib prints Matplotlib is building the font cache; this may take a moment, that is normal first-run plotting setup. It is not a parser or collection error.")
    guide.heading("5.3 What your current numbers mean", 2, "current_numbers")
    guide.paragraph(
        "Your terminal output and screenshot report 89,277 normalized observations, 0 parser notes, 3 daily summary rows, 44,746 persistence rows and 2 overlap rows. These values are a one-day baseline/smoke test, not the final 30-day findings."
    )
    guide.table(
        ["Source", "Date", "Unique indicators", "Raw rows"],
        [
            ("FireHOL", "2026-09-16", "17,297", "34,594"),
            ("Spamhaus DROP", "2026-09-16", "1,724", "3,450"),
            ("blocklist.de", "2026-09-16", "25,725", "51,233"),
        ],
        widths=[1.7, 1.3, 1.75, 1.55],
        font_size=9.1,
    )
    guide.paragraph("Three daily rows means three source/day records for the same calendar day, not three days. To answer persistence, turnover and historical-change questions, continue the daily collection at a fixed UTC time and build at least 30 distinct dates.")
    guide.heading("5.4 Run records and independent comparisons", 2, "run_records_comparisons")
    guide.paragraph("The complete daily command preserves more than the cumulative results. Each run_records/<run-id>/ folder contains a JSON manifest, a timestamped event timeline, the per-run collection log, parser notes, command stdout/stderr logs, a human-readable summary and a copy of the analysis produced at that moment. This prevents a later rerun from hiding what happened on an earlier day.")
    add_code(guide.doc, r"""# Show the newest run record.
ls -td run_records/* | head -n 1

# Read its summary and machine-readable manifest.
cat "$(ls -td run_records/* | head -n 1)/run_summary.md"
cat "$(ls -td run_records/* | head -n 1)/run_manifest.json"

# List independent comparisons without changing results/.
find comparisons -maxdepth 3 -type f -print | sort""")
    guide.paragraph("On the first distinct collection day, comparison_status=not_ready is expected. On the next distinct day, the program creates comparisons/<run-id>/<previous-day>_vs_<current-day>/. Use snapshot_index.csv to identify the exact input snapshots, source_summary.csv for counts and indicator_changes.csv or the source-specific *_changes.csv file for individual differences. If a source snapshot is missing on one selected day, its change counts are marked NA rather than treated as removals. The comparison folder is separate from results/ and never overwrites it.")
    if SCREENSHOT.exists():
        guide.figure(SCREENSHOT, "Figure 1", "The terminal evidence from the completed collection, normalization and analysis commands.", width=6.35)

    guide.heading("6. Add Danish evidence and classify candidates", 1, "danish_classification")
    guide.paragraph("Member 4 owns this stage. The current collector does not call a country or reputation API. Member 4 first retrieves permitted passive evidence, records its source and timestamp in CSV files, and then runs the transparent local classifier.")
    guide.heading("6.1 Create the evidence files", 2, "create_evidence_files")
    add_code(guide.doc, r"""cp data/danish_evidence_template.csv data/danish_evidence.csv
cp data/danish_prefixes_template.csv data/danish_prefixes.csv""")
    guide.paragraph("Open the copied files in VS Code if the code command is installed:")
    add_code(guide.doc, r"""code data/danish_evidence.csv
code data/danish_prefixes.csv""")
    guide.paragraph("On macOS, if code is not available, open them from Finder/VS Code or use:")
    add_code(guide.doc, r"""open -a "Visual Studio Code" data/danish_evidence.csv
open -a "Visual Studio Code" data/danish_prefixes.csv""")
    guide.paragraph("The downloaded project is a folder of scripts and data. VS Code is an editor for opening the folder and files; the commands themselves run in Terminal inside the activated .venv.")
    guide.heading("6.2 Run the classifier", 2, "classify_command")
    add_code(guide.doc, r"""python scripts/classify_danish.py \
  --input data/normalized/observations.csv \
  --evidence data/danish_evidence.csv \
  --prefixes data/danish_prefixes.csv \
  --out results/danish_classification.csv

column -s, -t < results/danish_classification.csv | head -n 20""")
    guide.paragraph("The four labels are DK-supported (at least two independent Danish signals), DK-weak (one signal), Conflicting (Danish and non-Danish signals disagree) and Unknown (no usable evidence). Do not infer Danish status from a .dk suffix alone, and do not treat any lookup as ground truth without recording its source and retrieval time.")
    guide.heading("6.3 Keep API use small and documented", 2, "api_guidance")
    guide.paragraph("If the group later uses RIPEstat, RIR/RDAP, DNS, AbuseIPDB, GreyNoise or VirusTotal, that is enrichment work—not part of the current three-feed collector. Agree on the candidate sample, access route, rate limits and retention rules first. Store API keys only in environment variables, never in CSV files, Git or the report. The report must state which service was used, for which fields, on which date and with what uncertainty.")

    guide.heading("7. Daily routine for 30 days", 1, "daily_routine")
    guide.paragraph("Run this sequence once per day at the same agreed UTC time. Do not run it repeatedly just to increase the row count. A second run on the same calendar date is a duplicate day for the primary daily comparison, although the raw file should remain for audit purposes.")
    add_code(guide.doc, r"""cd /absolute/path/to/Blocklist_Danish_Project_Package
source .venv/bin/activate
export BLOCKLIST_USER_AGENT='AAU-blocklist-study/0.1 contact: YOUR-GROUP-EMAIL'

# One command for the complete daily run.
python scripts/run_daily_pipeline.py

# The run folder contains the timestamped handoff.
ls -td run_records/* | head -n 1
cat "$(ls -td run_records/* | head -n 1)/run_summary.md"

# A comparison is created automatically after two distinct dates.
find comparisons -maxdepth 3 -type f -print | sort""")
    guide.heading("7.1 Daily status message to the group", 2, "daily_status")
    guide.paragraph("After each run, send a short message containing: collection date and UTC time; FireHOL/Spamhaus/blocklist.de status; whether URLhaus is disabled; parser-note count; whether a new date appeared; any duplicate or missing source; and the location of the updated raw snapshot and results files. This message becomes useful evidence for the methodology and contribution log.")
    guide.heading("7.2 How to count collection days", 2, "count_days")
    add_code(guide.doc, r"""cut -d, -f3 data/normalized/observations.csv | cut -c1-10 | sort -u | wc -l""")
    guide.paragraph("The expected final result is at least 30 distinct dates. If the count is lower, the historical study is incomplete even if the normalized row count is large.")

    guide.heading("8. Validation and final handoff", 1, "validation_handoff")
    guide.paragraph("Use these checks before giving the data to Member 5 or writing the final Findings chapter.")
    guide.heading("8.1 Check the scripts and preserve hashes", 2, "script_validation")
    add_code(guide.doc, r"""bash -n scripts/collect_snapshots.sh
python -m py_compile scripts/normalize_snapshot.py scripts/analyze_blocklists.py scripts/compare_snapshots.py scripts/run_daily_pipeline.py scripts/classify_danish.py
find data/raw -type f -exec shasum -a 256 {} + | sort > logs/raw_hash_register.txt""")
    guide.paragraph("bash -n checks shell syntax without downloading anything. py_compile checks Python syntax. The hash register provides a compact record of the raw evidence on macOS; on Linux, sha256sum can be used instead of shasum -a 256.")
    guide.heading("8.2 Rebuild the derived outputs", 2, "rebuild_outputs")
    add_code(guide.doc, r"""python scripts/normalize_snapshot.py \
  --raw-dir data/raw \
  --out data/normalized/observations.csv \
  --errors logs/normalization_errors.csv

python scripts/analyze_blocklists.py \
  --input data/normalized/observations.csv \
  --out-dir results

python scripts/compare_snapshots.py \
  --input data/normalized/observations.csv \
  --out-dir comparisons""")
    guide.heading("8.3 Final evidence checklist", 2, "final_checklist")
    for item in [
        "Every collection date has a log entry and a raw file or a documented failure.",
        "Every raw file has a source URL, retrieval timestamp, HTTP status, byte count and SHA-256 hash.",
        "The parser log was reviewed and all format changes are documented.",
        "The normalized file is regenerated from raw evidence, not edited by hand.",
        "Danish labels have evidence sources, timestamps and uncertainty notes.",
        "Every chart has a title, axis labels, caption, source description and collection window.",
        "Every numerical finding can be traced to a CSV result and the underlying snapshots.",
        "The six-person contribution log and final report cross-review are complete.",
    ]:
        guide.bullet(item)

    guide.heading("9. Safe practice commands", 1, "safe_practice")
    guide.paragraph("The project does not require malicious traffic, false abuse reports, brute force, exploit payloads or scans of public blocklist entries. If the group needs practice with network tools, use an isolated lab VM or a host that the group owns and has explicitly authorized.")
    add_code(guide.doc, r"""export TARGET=192.168.56.101
ip route
ping -c 2 $TARGET
nmap -Pn -sT --top-ports 100 --reason $TARGET
nmap -Pn -sV --version-light $TARGET""")
    guide.paragraph("Replace 192.168.56.101 only with the confirmed address of your own lab target. Do not substitute your Mac’s Wi-Fi IP 192.168.0.161 unless the group has explicitly decided that the Mac itself is the owned lab target and the supervisor permits the exercise; it is not needed for the blocklist study.")

    guide.heading("10. Optional scheduled collection", 1, "scheduled_collection")
    guide.paragraph("After the group has tested the manual routine and agreed on the collection time, cron can run the complete pipeline. Do not automate before the user agent, destination folder, retention plan and provider terms have been reviewed.")
    add_code(guide.doc, r"""crontab -e

15 03 * * * cd /absolute/path/to/Blocklist_Danish_Project_Package && BLOCKLIST_USER_AGENT='AAU-blocklist-study/0.1 contact: YOUR-GROUP-EMAIL' .venv/bin/python scripts/run_daily_pipeline.py >> logs/cron.log 2>&1""")
    guide.paragraph("The example runs at 03:15 according to the machine’s cron time zone. Confirm the machine time zone before relying on it. Inspect run_records/ after each automated execution; a cron entry alone is not evidence that a successful snapshot was obtained.")

    guide.heading("11. What to do next", 1, "next_action")
    guide.paragraph("You have already completed a successful baseline: the three enabled feeds were downloaded, normalization produced 89,277 observations and zero parser notes, and the analysis produced the first daily, persistence and overlap tables. Your next action is to agree with the group on one daily UTC collection time, use the complete pipeline and continue until at least 30 distinct dates are present.")
    guide.paragraph("Your hands-on handoff is: Member 3 maintains raw snapshots, hashes, logs, normalization and repeatable analysis; Member 4 adds the Danish evidence table and classification; Member 5 turns the stable outputs into final visualizations; Members 1, 2 and 6 connect the measurements to the introduction, literature, policy discussion and conclusion.")
    guide.heading("11.1 One-minute explanation for the group", 2, "one_minute_explanation")
    add_code(guide.doc, r"""We collect public FireHOL, Spamhaus DROP and blocklist.de snapshots once per day.
The scripts preserve the original files, normalize different formats into one table,
classify Danish-related evidence separately, and calculate size, additions,
removals, persistence and overlap. We compare the providers' documented policies
and delisting procedures. A blocklist entry is an observation to analyze, not proof
of current maliciousness. We need at least 30 distinct dates before final findings.""")

    guide.heading("12. Command index", 1, "command_index")
    guide.table(
        ["Task", "Owner", "Main command"],
        [
            ("Open/check project", "Everyone", "pwd; ls; find ..."),
            ("Create environment", "Everyone once", "python3 -m venv .venv"),
            ("Activate environment", "Everyone per session", "source .venv/bin/activate"),
            ("Install packages", "Everyone once", "python -m pip install -r requirements.txt"),
            ("Complete daily run", "Member 3", "python scripts/run_daily_pipeline.py"),
            ("Collect snapshots manually", "Member 3", "bash scripts/collect_snapshots.sh"),
            ("Review collection", "Member 3", "column -s, -t < logs/collection_log.csv"),
            ("Normalize", "Member 3", "python scripts/normalize_snapshot.py ..."),
            ("Review parser", "Member 3", "cat logs/normalization_errors.csv"),
            ("Analyze", "Member 5 / Member 3 run", "python scripts/analyze_blocklists.py ..."),
            ("Compare two dates", "Member 3 / Member 5", "python scripts/compare_snapshots.py ..."),
            ("Classify Danish evidence", "Member 4", "python scripts/classify_danish.py ..."),
            ("Validate", "Member 3", "bash -n; python -m py_compile; shasum"),
            ("Safe lab practice", "Approved lab only", "nmap commands in Section 9"),
        ],
        widths=[1.75, 1.45, 3.0],
        font_size=8.8,
    )
    guide.paragraph("For the full academic explanation, role division, theory, literature, findings structure and report roadmap, use the companion Member 3 and 4 Hands-on Project Guide.")

    toc_entries = [(title, level, anchor) for title, level, anchor in guide.headings if title != "Table of Contents"]
    replace_placeholder_with_links(guide.doc, "[[TOC]]", toc_entries)

    output = REPORT_DIR / "Blocklist_Project_Command_Runbook.docx"
    guide.doc.save(output)
    print(output)
    return output


if __name__ == "__main__":
    build_runbook()
