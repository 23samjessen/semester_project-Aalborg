from __future__ import annotations

import csv
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from make_member_3_4_guide import (
    FIG_DIR,
    REPORT_DIR,
    Guide,
    add_code,
    add_page_field,
    replace_placeholder_with_links,
)


ROOT = Path(__file__).resolve().parent


def set_footer(guide: Guide) -> None:
    footer = guide.doc.sections[0].footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Colleague Project Explainer  |  ")
    run.font.name = "Liberation Serif"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(90, 90, 90)
    add_page_field(footer)


def add_cover(guide: Guide) -> None:
    doc = guide.doc
    title_style = doc.styles["Title"]
    title_style_ppr = title_style._element.get_or_add_pPr()
    for child in list(title_style_ppr):
        if child.tag == qn("w:pBdr"):
            title_style_ppr.remove(child)
    p = doc.add_paragraph(style="Title")
    paragraph_ppr = p._p.get_or_add_pPr()
    for child in list(paragraph_ppr):
        if child.tag == qn("w:pBdr"):
            paragraph_ppr.remove(child)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    run = p.add_run("Our Blocklist Project — Easy Explanation")
    run.bold = True
    run.font.name = "Liberation Serif"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0, 0, 0)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p2.add_run("Analysis of Public Blocklists from a Danish Perspective")
    run.bold = True
    run.font.name = "Liberation Serif"
    run.font.size = Pt(15)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(25)
    run = p3.add_run("A short briefing for our six-person group")
    run.italic = True
    run.font.name = "Liberation Serif"
    run.font.size = Pt(11.5)

    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(90)
    for line in [
        "Aalborg University – Cyber Security",
        "Version: 16 September 2026",
        "Replace member names before sharing with the group",
    ]:
        run = p4.add_run(line + "\n")
        run.font.name = "Liberation Serif"
        run.font.size = Pt(11)

    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p5.paragraph_format.space_before = Pt(85)
    run = p5.add_run("The purpose of this document is to make the plan understandable and easy to coordinate.")
    run.font.name = "Liberation Serif"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(90, 90, 90)


def current_baseline() -> dict[str, object]:
    values: dict[str, object] = {
        "normalized": "not generated",
        "parser_notes": "not generated",
        "daily": "not generated",
        "persistence": "not generated",
        "overlap": "not generated",
        "dates": [],
    }
    normalized = ROOT / "data" / "normalized" / "observations.csv"
    if normalized.exists():
        with normalized.open(newline="", encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle))
        values["normalized"] = len(rows)
        values["dates"] = sorted({row.get("retrieved_at_utc", "")[:10] for row in rows if row.get("retrieved_at_utc")})
    errors = ROOT / "logs" / "normalization_errors.csv"
    if errors.exists():
        with errors.open(newline="", encoding="utf-8") as handle:
            values["parser_notes"] = sum(1 for _ in csv.DictReader(handle))
    for key, filename in (("daily", "daily_summary.csv"), ("persistence", "persistence.csv"), ("overlap", "overlap.csv")):
        path = ROOT / "results" / filename
        if path.exists():
            with path.open(newline="", encoding="utf-8") as handle:
                values[key] = sum(1 for _ in csv.DictReader(handle))
    return values


def build_colleague_explainer() -> Path:
    values = current_baseline()
    guide = Guide()
    set_footer(guide)
    add_cover(guide)

    guide.heading("Table of Contents", 1, "toc")
    guide.paragraph("Every entry below is linked. Click a heading to jump to that explanation.")
    guide.doc.add_paragraph("[[TOC]]")

    guide.heading("1. Our project in plain language", 1, "project_plain_language")
    guide.paragraph("We are studying public cybersecurity blocklists. A blocklist is a published list of IP addresses, network ranges, domains or URLs that a provider considers risky or associated with unwanted activity. Defenders may use these lists to block traffic or create alerts.")
    guide.paragraph("Our group wants to understand how these lists work in practice: how they publish entries, how the entries change, how long entries remain visible, how much two providers agree, and what evidence can support a Danish connection.")
    guide.paragraph("The key point is that a list entry is an observation, not automatic proof of current criminal activity. We will report what the data and the provider documentation support, including uncertainty.")
    guide.heading("1.1 The question we are answering", 2, "simple_research_question")
    guide.paragraph("How do selected public blocklists change over time, and what can we learn about indicators associated with Denmark from a careful 30-day observation period?")
    guide.heading("1.2 What success looks like", 2, "success_definition")
    guide.table(
        ["Part", "A successful result"],
        [
            ("Theory", "A clear explanation of blocklists, false positives, attribution, measurement and source criticism."),
            ("Data", "Dated original snapshots with URLs, timestamps, status, size and SHA-256 hashes."),
            ("Analysis", "Reproducible CSV tables for counts, additions, removals, persistence and overlap."),
            ("Danish perspective", "Evidence-based labels with source, timestamp and uncertainty; no unsupported country claim."),
            ("Report", "One connected academic document with linked contents, figures, tables, IEEE references, index and appendices."),
        ],
        widths=[1.6, 5.0],
        font_size=9.3,
    )

    guide.heading("2. What are we collecting?", 1, "what_collecting")
    guide.paragraph("The collector downloads public feed files over HTTPS. It saves the original response in a dated folder so we can later show exactly what was available on that day.")
    guide.table(
        ["Source", "What it gives us", "Current status"],
        [
            ("FireHOL Level 2", "Public IP and CIDR entries from a documented aggregated list.", "Enabled"),
            ("Spamhaus DROP", "Public DROP IPv4 entries and provider fields when supplied.", "Enabled"),
            ("blocklist.de", "Public text entries representing reported attack-source IP addresses.", "Enabled"),
            ("URLhaus", "Potential domain/URL feed for a separate scope.", "Disabled until access route and terms are reviewed"),
        ],
        widths=[1.5, 3.65, 1.45],
        font_size=9.0,
    )
    guide.heading("2.1 Is this an API?", 2, "api_explanation")
    guide.paragraph("In the current run, the program is not using a private API with a key. The Bash collector uses curl to request three public HTTPS feed URLs. The response is saved as a file and logged. The installed Python package `requests` is available for future controlled enrichment, but it is not currently calling AbuseIPDB, GreyNoise, VirusTotal or RIPEstat.")
    guide.paragraph("This distinction matters when we explain the project: we are collecting published public feeds, not accessing private provider databases. If we later add an API, we must document its endpoint, purpose, access terms, rate limit, timestamp and key storage.")
    guide.heading("2.2 What the program does not do", 2, "not_active_scanning")
    for item in [
        "It does not scan the Mac Wi-Fi address 192.168.0.161.",
        "It does not connect to, test or exploit the IP addresses found in a feed.",
        "It does not send false abuse reports or delisting requests.",
        "It does not decide that an indicator is malicious just because it appears once.",
        "It does not decide that an indicator is Danish without recorded evidence.",
    ]:
        guide.bullet(item)

    guide.heading("3. The theoretical part", 1, "theory_part")
    guide.paragraph("The report must explain the ideas behind the measurements before showing the numbers. This is the part mainly led by Members 1, 2 and 6, with technical input from Members 3, 4 and 5.")
    guide.table(
        ["Topic", "Simple explanation", "Where it appears"],
        [
            ("Blocklist types", "Different providers have different purposes and update rules.", "Background and technical context"),
            ("False positives", "A benign owner can be affected by an entry that is old, broad or no longer accurate.", "Problem and discussion"),
            ("Danish attribution", "Country allocation, geolocation, ASN and DNS are different signals and can disagree.", "Methodology and limitations"),
            ("Longitudinal measurement", "Repeated snapshots let us measure change instead of describing one moment.", "Methodology"),
            ("Source criticism", "A provider list is useful evidence, but it has scope, bias and uncertainty.", "Literature and discussion"),
        ],
        widths=[1.55, 3.45, 1.6],
        font_size=9.0,
    )
    guide.paragraph("The literature review should connect these ideas to peer-reviewed work, RFCs and official provider documentation. Member 2 maintains the IEEE reference list and search log so the report can distinguish published knowledge from our own measurements.")

    guide.heading("4. The hands-on part", 1, "hands_on_part")
    guide.paragraph("The hands-on part is a repeatable sequence. Member 3 owns the collection and normalization routine; Member 4 adds Danish evidence; Member 5 analyzes the stable data. The other members use the outputs in the report.")
    pipeline = FIG_DIR / "figure_1_research_pipeline.png"
    if pipeline.exists():
        guide.figure(pipeline, "Figure 1", "The complete path from public source files to report interpretation.", width=6.25)
    guide.heading("4.1 What happens after one command", 2, "one_command_explanation")
    guide.table(
        ["Stage", "What happens", "Output"],
        [
            ("Collect", "curl requests the enabled public URLs and saves the original files.", "data/raw/YYYY-MM-DD/ and collection_log.csv"),
            ("Normalize", "Python reads each supported format, validates IP/CIDR values and keeps provenance.", "data/normalized/observations.csv"),
            ("Analyze", "Python groups by source and day, then compares sets of canonical values.", "results/*.csv and PNG figures"),
            ("Record", "The daily runner stores UTC times, commands, statuses, hashes and a copy of the analysis.", "run_records/<run-id>/"),
            ("Compare", "The program compares the latest snapshots from two distinct dates.", "comparisons/<run-id>/"),
            ("Classify", "Evidence supplied by Member 4 is combined into a cautious Danish label.", "results/danish_classification.csv"),
            ("Write", "Verified tables and figures are connected to the academic report sections.", "report/*.docx"),
        ],
        widths=[1.05, 3.7, 1.85],
        font_size=8.9,
    )
    add_code(guide.doc, "source .venv/bin/activate\nexport BLOCKLIST_USER_AGENT='AAU-blocklist-study/0.1 contact: GROUP-EMAIL'\npython scripts/run_daily_pipeline.py")
    guide.paragraph("The root `README.md` contains the complete copy-paste command list with a comment above every command. The `Blocklist_Project_Command_Runbook.docx` contains the same workflow with expected outputs and troubleshooting. The older collector, normalizer and analyzer commands remain available for troubleshooting one stage at a time.")
    guide.heading("4.2 Where the daily record is saved", 2, "daily_record_and_comparison")
    guide.paragraph("Every complete daily run creates a unique `run_records/<run-id>/` folder. It contains `run_manifest.json`, `run_events.csv`, the per-run `collection.csv`, command output logs, parser notes, a human-readable `run_summary.md` and a copy of the cumulative analysis. This means the group can reconstruct what happened on a particular day even after the main results are regenerated.")
    guide.paragraph("After at least two distinct dates, the runner creates `comparisons/<run-id>/<previous-day>_vs_<current-day>/`. It contains a snapshot index, source summary, individual added and removed indicators, one change file per source, notes and a figure. If a provider is missing on one selected day, its source summary is marked as missing and uses NA for change counts instead of falsely reporting removals. The comparison folder is separate from `results/`, so an error in the comparison cannot change the main analysis.")

    guide.heading("5. Why do we need 30 days?", 1, "why_thirty_days")
    guide.paragraph("The first run is a baseline. It tells us that the pipeline works, but it cannot tell us how an indicator behaves over time. Three source rows from one date are not the same as three days of history.")
    timeline = FIG_DIR / "figure_4_collection_timeline.png"
    if timeline.exists():
        guide.figure(timeline, "Figure 2", "A simple 30-day schedule with quality checks during the collection period.", width=6.25)
    guide.paragraph("With at least 30 distinct dates, we can compare a source with its preceding successful snapshot. This supports measured statements about additions, removals, persistence, turnover and overlap across time. We can also show missing or failed days instead of mistaking them for removals.")
    guide.table(
        ["If we collect...", "What we can honestly say"],
        [
            ("One date", "How large each source was on that date and how much exact overlap was visible then."),
            ("Several dates", "Whether counts and membership changed between the dates we successfully captured."),
            ("At least 30 dates", "A stronger descriptive picture of persistence and turnover during the defined collection window."),
        ],
        widths=[1.65, 4.95],
        font_size=9.3,
    )

    guide.heading("6. How we divide the work", 1, "work_division")
    guide.paragraph("Each member owns one part, but the final report is one project. Everyone should understand the complete pipeline well enough to explain how another member’s output was produced.")
    guide.table(
        ["Member", "Main job", "Must hand over"],
        [
            ("1 — Coordinator", "Problem, aim, research questions, meetings and final integration.", "Clean report structure, decisions log and final consistency check."),
            ("2 — Literature", "Theory, literature review, IEEE references and source register.", "Search log, annotated sources and cited theory for each research question."),
            ("3 — Collection", "Run the collector, preserve raw feeds, check hashes, normalize and record failures.", "Dated raw snapshots, collection log, normalized CSV, parser log and run notes."),
            ("4 — Danish evidence", "Collect allocation/geolocation/ASN/DNS evidence and apply cautious labels.", "Evidence CSV, prefix CSV, source/timestamp notes and classification output."),
            ("5 — Analysis", "Check the analysis tables and create final visualizations.", "Verified statistics, figures, captions and links to source CSV files."),
            ("6 — Policy and conclusion", "Compare provider policies, ethics, correction/delisting and limitations.", "Policy matrix, discussion draft, limitation text and conclusion linked to findings."),
        ],
        widths=[1.35, 2.65, 2.6],
        font_size=8.55,
    )
    guide.paragraph("Our hands-on work belongs mostly to Members 3 and 4, but it feeds the analysis and report. If you take the hands-on role, your job is not just to run commands: you must preserve evidence, report failures, explain the data model and give the group a reproducible handoff.")

    guide.heading("7. Who waits for whom?", 1, "dependencies")
    guide.paragraph("Some tasks can start immediately. Other tasks need a reliable input from an earlier stage. This is normal project coordination, not a problem.")
    guide.table(
        ["Member", "Can start now", "Final part waits for"],
        [
            ("3", "Set up the environment, test the parser, collect the baseline and schedule the daily run.", "The group’s agreed source list and collection time."),
            ("4", "Read provider documentation and design the evidence sheet.", "The normalized candidate indicators from Member 3 for final classification."),
            ("5", "Prepare figure templates, definitions and quality checks.", "Enough collection dates and the reviewed Danish classification."),
            ("6", "Draft provider policy, ethics and safe correction procedures.", "Verified findings for the final discussion and conclusion."),
            ("1 and 2", "Write structure, problem statement, theory and literature review.", "Final numbers, figures and conclusions for the integrated report."),
        ],
        widths=[1.05, 3.1, 2.45],
        font_size=8.9,
    )
    guide.paragraph("The practical order is: Member 3 produces a clean observation table; Member 4 adds country evidence; Member 5 measures and visualizes; Member 6 interprets policies and limitations; Member 1 integrates; Member 2 keeps every theoretical claim and reference supported.")

    guide.heading("8. What we have already done", 1, "current_status")
    date_text = ", ".join(values["dates"]) if values["dates"] else "no date recorded"
    guide.paragraph(f"The current package contains a successful baseline run. It has {values['normalized']:,} normalized observations, {values['parser_notes']} parser notes, {values['daily']} daily summary rows, {values['persistence']:,} persistence rows and {values['overlap']} overlap rows. The date(s) currently present are: {date_text}.")
    guide.paragraph("Your reported Mac terminal run showed 89,277 normalized observations and 44,746 persistence rows. The bundled sample folder can show a different count if it contains one snapshot per source while your Mac data contains two snapshots from separate runs on the same date. The count is determined by the raw files under data/raw; it is not typed into the analysis. For the final report, keep the exact raw snapshot set and regenerated numbers used by the group.")
    guide.paragraph("This is enough to demonstrate that the program runs and that the three enabled feeds were parsed. It is not enough to write a final 30-day persistence conclusion. We keep collecting instead of pretending that the baseline is the completed study.")
    guide.table(
        ["File or folder", "What it proves"],
        [
            ("data/raw/", "What each public source returned at a specific time."),
            ("logs/collection_log.csv", "Whether a download succeeded and its URL, size and hash."),
            ("data/normalized/observations.csv", "How the different feed shapes were represented in one schema."),
            ("logs/normalization_errors.csv", "Parser notes; the header only means zero notes were recorded."),
            ("results/", "Derived descriptive tables and charts; not raw evidence."),
            ("run_records/", "Timestamped audit records for each complete daily run."),
            ("comparisons/", "Independent comparisons between two distinct dates."),
        ],
        widths=[2.35, 4.25],
        font_size=9.1,
    )

    guide.heading("9. A simple 30-day working plan", 1, "thirty_day_plan")
    guide.table(
        ["Period", "Team action", "Checkpoint"],
        [
            ("Before Day 1", "Agree feeds, UTC collection time, evidence rules, file naming and review owners.", "Supervisor-approved scope and contact string."),
            ("Days 1–3", "Run the baseline, test normalization, check source semantics and fix obvious parser issues.", "All enabled feeds parse; failures are logged."),
            ("Days 4–14", "Collect daily, review logs, begin Danish evidence and keep the literature/source register updated.", "No unexplained gaps; weekly backup and review."),
            ("Days 15–21", "Run a midpoint analysis, check definitions and compare the first trends with provider policies.", "Group review of preliminary tables and limitations."),
            ("Days 22–30", "Continue the same routine; do not change the method without recording the change.", "At least 30 distinct dates or a documented reason why not."),
            ("After Day 30", "Freeze the dataset, regenerate outputs, verify figures/references and write findings/discussion/conclusion.", "Final evidence package and integrated report."),
        ],
        widths=[1.25, 3.55, 1.8],
        font_size=8.7,
    )

    guide.heading("10. What to say in the group meeting", 1, "meeting_script")
    guide.paragraph("We are running a controlled observational study. We collect public FireHOL, Spamhaus DROP and blocklist.de snapshots once per day, save the originals and their hashes, normalize them into one table, and calculate descriptive changes. Danish status is classified separately from list membership. We need at least 30 distinct dates because one snapshot cannot measure persistence or turnover. Each member owns a work package, and the handoffs are clear: collection, Danish evidence, analysis, policy interpretation and final integration.")
    guide.paragraph("If someone asks whether the program is connected to an API, the accurate short answer is: the current collector uses public HTTPS feed URLs with curl, not an authenticated API key. URLhaus is disabled. Optional enrichment services are not being queried unless we add and document that step.")
    guide.paragraph("If someone asks whether the current result is final, the accurate answer is: no. The baseline proves that the pipeline runs. Final time-based findings wait for the agreed collection window and the group’s review.")

    guide.heading("11. What each member should deliver", 1, "deliverables")
    guide.table(
        ["Member", "Small deliverable for the shared folder"],
        [
            ("1", "One-page problem statement, decision log and final report checklist."),
            ("2", "Literature search table with IEEE entries and one paragraph explaining the research gap."),
            ("3", "Daily collection log, raw-file hash register, normalization command output and data dictionary notes."),
            ("4", "Danish evidence and prefix files with source, access date, signal type and uncertainty."),
            ("5", "Analysis tables, three final figures, captions and a short explanation of each metric."),
            ("6", "Provider policy comparison, safe correction/delisting discussion, limitations and conclusion draft."),
        ],
        widths=[1.0, 5.6],
        font_size=9.15,
    )
    guide.paragraph("Use clear file names and do not overwrite raw snapshots. Put questions or failed runs in the shared project log so the next member can continue without guessing.")

    guide.heading("12. What we do next", 1, "next_steps")
    for item in [
        "Agree the daily UTC collection time and the final 30-day window with the supervisor.",
        "Replace the example user-agent contact with a group-controlled email address.",
        "Run the complete daily pipeline once per day and review the newest run_records/<run-id>/run_summary.md.",
        "Keep URLhaus disabled unless the group approves its route and terms.",
        "Let Member 4 build evidence files from documented sources rather than guessing from a domain suffix.",
        "Let Member 5 use only regenerated result files for final charts.",
        "Hold a short weekly review: days collected, failed sources, parser notes, Danish evidence, open decisions and report progress.",
    ]:
        guide.numbered(item)
    add_code(guide.doc, "# Run from the project root\nsource .venv/bin/activate\nexport BLOCKLIST_USER_AGENT='AAU-blocklist-study/0.1 contact: GROUP-EMAIL'\npython scripts/run_daily_pipeline.py\n\n# Inspect the newest run record\nls -td run_records/* | head -n 1\ncat \"$(ls -td run_records/* | head -n 1)/run_summary.md\"")
    guide.paragraph("For the full commands, code explanation, Arabic summary, supervisor plan and detailed hands-on guide, use the other files included in this package.")

    guide.heading("13. Project control rules", 1, "control_rules")
    for item in [
        "One person runs the daily collection; another person checks the log at least weekly.",
        "Raw files are evidence and are never edited by hand.",
        "Each complete run has its own timestamped manifest, event log and archived analysis.",
        "Comparisons are stored outside `results/` and can be checked independently.",
        "Every derived table can be regenerated from the raw files.",
        "Every Danish label has an evidence source, timestamp and uncertainty note.",
        "Every result in the report points to a table, figure, raw snapshot or cited policy source.",
        "The group reports missing data and limitations openly; a missing day or source snapshot is not silently treated as an indicator removal.",
    ]:
        guide.bullet(item)

    toc_entries = [(title, level, anchor) for title, level, anchor in guide.headings if title != "Table of Contents"]
    replace_placeholder_with_links(guide.doc, "[[TOC]]", toc_entries)

    output = REPORT_DIR / "Project_Colleague_Explainer.docx"
    guide.doc.save(output)
    print(output)
    return output


if __name__ == "__main__":
    build_colleague_explainer()
