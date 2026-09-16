from __future__ import annotations

import re
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
REPORT_DIR = ROOT / "report"
FIG_DIR = ROOT / "figures"
SCREENSHOT = ROOT.parent / "upload" / "de44a4ec-28a8-4d16-849b-3ecdd384295d.png"
REPORT_DIR.mkdir(parents=True, exist_ok=True)
FIG_DIR.mkdir(parents=True, exist_ok=True)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "B7C5D4", size: str = "6") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_hyperlink(paragraph, label: str, anchor: str, color: str = "1F4E79") -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), color)
    rpr.append(colour)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def replace_placeholder_with_links(doc: Document, token: str, entries: list[tuple[str, int, str]]) -> None:
    placeholder = None
    for paragraph in doc.paragraphs:
        if token in paragraph.text:
            placeholder = paragraph
            break
    if placeholder is None:
        return
    parent = placeholder._p.getparent()
    index = parent.index(placeholder._p)
    parent.remove(placeholder._p)
    for offset, (label, level, anchor) in enumerate(entries):
        p = OxmlElement("w:p")
        p_pr = OxmlElement("w:pPr")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        p_pr.append(spacing)
        if level > 1:
            indent = OxmlElement("w:ind")
            indent.set(qn("w:left"), str((level - 1) * 360))
            p_pr.append(indent)
        p.append(p_pr)
        parent.insert(index + offset, p)
        paragraph = docx_paragraph_from_element(p, parent)
        add_internal_hyperlink(paragraph, label, anchor)


def docx_paragraph_from_element(element, parent):
    from docx.text.paragraph import Paragraph

    return Paragraph(element, parent)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(end)


def add_code(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.right_indent = Inches(0.22)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.05
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F4F7")
    p_pr.append(shd)
    border = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement("w:" + edge)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "4")
        element.set(qn("w:color"), "D5DCE3")
        border.append(element)
    p_pr.append(border)
    run = paragraph.add_run(text)
    run.font.name = "Liberation Mono"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Mono")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Mono")
    run.font.size = Pt(9.2)


def add_caption(doc: Document, label: str, text: str) -> None:
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{label}. {text}")
    r.italic = True


class Guide:
    def __init__(self) -> None:
        self.doc = Document()
        self.headings: list[tuple[str, int, str]] = []
        self.bookmark_id = 1
        self.figures: list[tuple[str, str]] = []
        self.configure()

    def configure(self) -> None:
        section = self.doc.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.79)
        section.right_margin = Inches(0.79)
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)

        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = "Liberation Serif"
        normal._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Serif")
        normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Serif")
        normal.font.size = Pt(11.2)
        normal.paragraph_format.line_spacing = 1.18
        normal.paragraph_format.space_after = Pt(8)

        for style_name, size, before, after in [
            ("Heading 1", 16, 18, 8),
            ("Heading 2", 12.8, 12, 5),
            ("Heading 3", 11.2, 8, 3),
        ]:
            style = styles[style_name]
            style.font.name = "Liberation Serif"
            style._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Serif")
            style._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Serif")
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True
            if style_name == "Heading 1":
                style.paragraph_format.page_break_before = True

        caption = styles["Caption"]
        caption.font.name = "Liberation Serif"
        caption.font.size = Pt(9.5)
        caption.font.italic = True
        caption.font.color.rgb = RGBColor(70, 70, 70)

        for style_name in ("List Bullet", "List Bullet 2", "List Number"):
            style = styles[style_name]
            style.font.name = "Liberation Serif"
            style._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Serif")
            style._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Serif")
            style.font.size = Pt(11.0)
            style.paragraph_format.line_spacing = 1.15
            style.paragraph_format.space_after = Pt(3)

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Member 3 and 4 Hands-on Guide  |  ")
        run.font.name = "Liberation Serif"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(90, 90, 90)
        add_page_field(footer)

    def heading(self, title: str, level: int = 1, anchor: str | None = None) -> str:
        paragraph = self.doc.add_paragraph(style=f"Heading {level}")
        paragraph.add_run(title)
        if anchor is None:
            anchor = re.sub(r"[^a-z0-9]+", "_", title.lower()).strip("_")
        add_bookmark(paragraph, anchor, self.bookmark_id)
        self.bookmark_id += 1
        self.headings.append((title, level, anchor))
        return anchor

    def paragraph(self, text: str, bold_lead: str | None = None) -> None:
        paragraph = self.doc.add_paragraph()
        if bold_lead and text.startswith(bold_lead):
            r = paragraph.add_run(bold_lead)
            r.bold = True
            paragraph.add_run(text[len(bold_lead):])
        else:
            paragraph.add_run(text)

    def bullet(self, text: str, level: int = 0) -> None:
        style = "List Bullet" if level == 0 else "List Bullet 2"
        self.doc.add_paragraph(text, style=style)

    def numbered(self, text: str) -> None:
        self.doc.add_paragraph(text, style="List Number")

    def table(self, headers: list[str], rows: list[tuple[str, ...]], widths: list[float] | None = None, font_size: float = 9.4) -> None:
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        header = table.rows[0]
        set_repeat_table_header(header)
        for index, value in enumerate(headers):
            cell = header.cells[index]
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "25577E")
            set_cell_border(cell)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Liberation Serif"
                    run.font.size = Pt(font_size)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
        for row_index, values in enumerate(rows):
            row = table.add_row()
            for col_index, value in enumerate(values):
                cell = row.cells[col_index]
                cell.text = value
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_shading(cell, "F4F7FA" if row_index % 2 else "FFFFFF")
                set_cell_border(cell)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.05
                    for run in paragraph.runs:
                        run.font.name = "Liberation Serif"
                        run.font.size = Pt(font_size)
        if widths:
            for row in table.rows:
                for cell, width in zip(row.cells, widths):
                    cell.width = Inches(width)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def figure(self, path: Path, label: str, caption: str, width: float = 6.25) -> None:
        self.doc.add_picture(str(path), width=Inches(width))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_caption(self.doc, label, caption)
        self.figures.append((label, caption))


def make_pipeline_figure() -> Path:
    path = FIG_DIR / "member_3_4_program_pipeline.png"
    fig, axis = plt.subplots(figsize=(11, 2.8))
    axis.set_xlim(0, 12)
    axis.set_ylim(0, 3)
    axis.axis("off")
    boxes = [
        (0.3, "Public HTTPS\nfeeds", "#DDEAF5"),
        (3.0, "data/raw/\noriginal snapshots", "#E7F1E8"),
        (5.8, "normalize\nobservations.csv", "#F7EBD2"),
        (8.7, "analyze\nresults + figures", "#E9E0F1"),
    ]
    for x, label, color in boxes:
        patch = FancyBboxPatch((x, 0.95), 2.05, 0.95, boxstyle="round,pad=0.03,rounding_size=0.06", facecolor=color, edgecolor="#25577E", linewidth=1.5)
        axis.add_patch(patch)
        axis.text(x + 1.025, 1.43, label, ha="center", va="center", fontsize=11, color="#1D2D3A")
    for x in (2.45, 5.25, 8.15):
        axis.annotate("", xy=(x + 0.45, 1.43), xytext=(x, 1.43), arrowprops={"arrowstyle": "->", "color": "#25577E", "linewidth": 1.8})
    axis.text(6, 2.52, "What the program does", ha="center", va="center", fontsize=16, fontweight="bold", color="#25577E")
    axis.text(1.3, 0.43, "not a scan", ha="center", va="center", fontsize=9, color="#666666")
    axis.text(4.0, 0.43, "evidence trail", ha="center", va="center", fontsize=9, color="#666666")
    axis.text(6.8, 0.43, "common format", ha="center", va="center", fontsize=9, color="#666666")
    axis.text(9.7, 0.43, "descriptive output", ha="center", va="center", fontsize=9, color="#666666")
    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return path


def add_cover(guide: Guide) -> None:
    doc = guide.doc
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    r = p.add_run("Member 3 and 4 Hands-on Guide")
    r.bold = True
    r.font.name = "Liberation Serif"
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(31, 78, 121)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("Analysis of Public Blocklists from a Danish Perspective")
    r.font.name = "Liberation Serif"
    r.font.size = Pt(15)
    r.bold = True
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(25)
    r = p3.add_run("What we are studying, what the program does, and how the six-person group will work")
    r.font.name = "Liberation Serif"
    r.font.size = Pt(11.5)
    r.font.italic = True
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(90)
    for line in [
        "Prepared for the project group",
        "Aalborg University – Cyber Security",
        "16 September 2026",
    ]:
        r = p4.add_run(line + "\n")
        r.font.name = "Liberation Serif"
        r.font.size = Pt(11)
    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p5.paragraph_format.space_before = Pt(85)
    r = p5.add_run("Use this guide as the group’s working explanation. The report’s final empirical findings must come from the completed collection period.")
    r.font.name = "Liberation Serif"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(90, 90, 90)


def build_guide() -> Path:
    pipeline = make_pipeline_figure()
    guide = Guide()
    add_cover(guide)

    guide.heading("Table of Contents", 1, "toc")
    guide.paragraph("The entries below are clickable. Select a heading to jump directly to that part of the guide.")
    guide.doc.add_paragraph("[[TOC]]")

    guide.heading("1. The project in one sentence", 1, "project_one_sentence")
    guide.paragraph("We are observing how public cybersecurity blocklists change over time, how their operators explain those changes, and what can be learned about indicators associated with Denmark.")
    guide.paragraph("The project is not trying to prove that every listed address is currently criminal or that every unlisted address is safe. A listing is a public data state. Our job is to measure that state, preserve evidence, compare sources and explain uncertainty.")
    guide.heading("1.1 What problem are we investigating?", 2, "problem_in_plain_language")
    guide.paragraph("Network defenders often use blocklists to reject traffic, increase a risk score or alert an analyst. The problem is that different lists have different purposes and update rules. An IP address can be reassigned, a domain can change owner or hosting, and a network range can cover more addresses than the original observation. A benign owner may therefore experience a false positive or may not know how to request correction.")
    guide.heading("1.2 The three research questions", 2, "research_questions_explained")
    guide.table(
        ["Question", "Meaning in simple words", "Evidence we need"],
        [
            ("RQ1", "How do the selected lists add, update, expire and remove indicators?", "Provider documentation, feed structure, timestamps and policy pages"),
            ("RQ2", "What listing, persistence, turnover and overlap patterns appear for Danish-related indicators?", "At least 30 days of timestamped snapshots and normalized observations"),
            ("RQ3", "How difficult is it for a benign owner to correct or remove an entry?", "Published removal procedures and, only if allowed, a controlled owned-asset case"),
        ],
        widths=[0.65, 3.05, 2.65],
        font_size=9.5,
    )
    guide.heading("1.3 What is the target?", 2, "target_of_study")
    guide.table(
        ["Item", "Is it a project target?", "Explanation"],
        [
            ("Public feed files", "Yes", "We download and preserve the published lists and their documented metadata."),
            ("IP addresses and CIDR ranges", "Yes", "We compare them as different representations and do not silently expand ranges."),
            ("Domains and URLs", "Yes, separately", "A domain or URL is not treated as an IP address; URLs can contain paths."),
            ("Danish-related evidence", "Yes, later", "We use separate allocation, geolocation, ASN, DNS and other evidence fields."),
            ("Your Mac: 192.168.0.161", "No", "This is your local Wi-Fi address. The collection program does not scan it."),
            ("Third-party public hosts", "No active probing", "A public listing is not permission to scan or connect to a host."),
        ],
        widths=[1.55, 1.15, 3.65],
        font_size=9.2,
    )

    guide.heading("2. Theory and hands-on work", 1, "theory_and_hands_on")
    guide.paragraph("The project has two connected sides. The theoretical side explains what blocklists mean and how earlier research describes their transparency and dynamics. The hands-on side collects real public feed states, converts them into a common structure and measures what changed. Neither side is enough by itself.")
    guide.table(
        ["Theoretical part", "Hands-on part", "Report location"],
        [
            ("Blocklists, DNSBLs, feeds, IPs, CIDR, domains and URLs", "Inspect the actual format of each selected source", "Chapter 2"),
            ("Transparency, persistence, turnover and false-positive risk", "Collect dated snapshots and calculate changes", "Chapters 3–5"),
            ("Country attribution and uncertainty", "Build Danish evidence and classify cases", "Chapters 2, 4 and 5"),
            ("Policy, correction and delisting", "Read and code removal procedures", "Chapters 3, 4, 5 and 6"),
            ("Scientific validity and reproducibility", "Keep raw files, hashes, logs, scripts and a data dictionary", "Chapter 4 and appendices"),
        ],
        widths=[2.25, 2.75, 1.35],
        font_size=9.2,
    )
    guide.heading("2.1 What the group must deliver", 2, "project_deliverables")
    for text in [
        "A clear problem statement, aim and research questions.",
        "A literature review with IEEE numbered references and a documented search process.",
        "A technical dissection of the selected feeds and their stated purpose.",
        "At least 30 consecutive days of raw snapshots, with collection status and file hashes.",
        "A normalized observation file that keeps source provenance and the original representation.",
        "A timestamped run record for every daily execution, including commands, statuses, hashes and archived analysis.",
        "An independent comparison folder for each pair of distinct collection dates.",
        "Danish classification evidence, including uncertainty and conflicting signals.",
        "Descriptive tables and figures for size, additions, removals, persistence and overlap.",
        "A policy comparison and a safe, evidence-based discussion of correction or delisting.",
        "A conclusion that answers RQ1–RQ3 from measured evidence, not assumptions.",
    ]:
        guide.bullet(text)

    guide.heading("3. The program you are running", 1, "program_explained")
    guide.paragraph("The program is a small local research pipeline. It downloads public feed files, stores the original snapshots, converts different formats into one observation table and produces descriptive results. It is not an attack tool and it is not a scanner.")
    guide.figure(pipeline, "Figure 1", "The program pipeline from public feed retrieval to normalized observations and descriptive results.")
    guide.heading("3.1 Which feeds and APIs are connected now?", 2, "current_feeds_and_apis")
    guide.paragraph("At the moment, the collector is connected to three public HTTPS feed URLs. These are direct file downloads, not authenticated API calls. The request method is a normal HTTP GET made by curl. No API key is used by the current successful run.")
    guide.table(
        ["Source", "Current route", "What it contributes", "Current status"],
        [
            ("FireHOL", "Raw GitHub `firehol_level2.netset` feed", "IP addresses and CIDR ranges from an aggregated public source", "Collected"),
            ("Spamhaus DROP", "Official DROP v4 JSON/JSON-lines feed", "Curated IPv4 network ranges with provider metadata", "Collected"),
            ("blocklist.de", "Official `all.txt` export", "Individual IP addresses from a short-retention attack-report list", "Collected"),
            ("URLhaus", "Optional URLhaus route", "Domains, hosts or URLs for the separate domain/URL dimension", "Disabled"),
        ],
        widths=[1.2, 2.2, 2.1, 0.9],
        font_size=8.9,
    )
    guide.paragraph("The report also mentions AbuseIPDB, GreyNoise, VirusTotal and RIPEstat. They are optional enrichment or evidence sources for later work; the current collector does not call them. APIs may require keys, subscriptions or fair-use checks. Member 4 should only use them for the small Danish candidate set and only after the group documents the access route.")
    guide.heading("3.2 What the commands did", 2, "commands_explained")
    guide.table(
        ["Command", "What it does", "Why we need it"],
        [
            ("`python3 -m venv .venv`", "Creates an isolated Python environment in `.venv`.", "Keeps this project’s packages separate from the rest of the Mac."),
            ("`source .venv/bin/activate`", "Activates that environment; the prompt shows `(.venv)`.", "Makes the following Python and pip commands use the project environment."),
            ("`python -m pip install -r requirements.txt`", "Installs the listed Python packages.", "Provides the libraries used for parsing, analysis, plotting and later enrichment."),
            ("`python scripts/run_daily_pipeline.py`", "Runs collection, normalization, analysis and comparison in order.", "Creates one timestamped audit record for the complete daily run."),
            ("`bash scripts/collect_snapshots.sh`", "Downloads the enabled feeds and records timestamps, status, size and SHA-256 hashes.", "Creates the raw evidence for one collection event."),
            ("`python scripts/normalize_snapshot.py ...`", "Reads raw files, parses their formats and writes one common CSV.", "Allows different sources to be compared without losing provenance."),
            ("`cat logs/normalization_errors.csv`", "Displays parser notes.", "A header with no rows means zero parser notes were recorded."),
            ("`python scripts/analyze_blocklists.py ...`", "Creates daily summaries, persistence, overlap and figures.", "Turns the normalized data into descriptive evidence for Chapters 5 and 6."),
            ("`python scripts/compare_snapshots.py ...`", "Compares the latest snapshots from two distinct days.", "Creates independent added/removed files without changing `results/`."),
        ],
        widths=[2.1, 2.55, 1.75],
        font_size=8.75,
    )
    guide.heading("3.3 What each folder means", 2, "folders_explained")
    guide.table(
        ["Location", "Meaning", "Rule"],
        [
            ("`data/raw/`", "Original files returned by providers.", "Never edit or overwrite them."),
            ("`logs/collection_log.csv`", "Retrieval time, URL, HTTP status, size, hash and status.", "Check after every collection."),
            ("`data/normalized/observations.csv`", "Derived common observation table.", "It can be regenerated from raw files."),
            ("`logs/normalization_errors.csv`", "Files or formats that need review.", "A blank data section means no parser notes."),
            ("`results/`", "Descriptive CSV tables and PNG figures.", "Do not call these final findings until coverage is complete."),
            ("`run_records/<run-id>/`", "One timestamped audit folder per complete daily run.", "Keep the manifest, events, command logs and archived analysis."),
            ("`comparisons/<run-id>/`", "Independent previous/current-day comparisons.", "It never overwrites `results/`."),
            ("`data/danish_*template.csv`", "Templates for Member 4’s country evidence.", "Fill with evidence actually retrieved by the group."),
        ],
        widths=[2.1, 2.75, 1.55],
        font_size=9.0,
    )

    guide.heading("4. Understanding your current terminal result", 1, "current_terminal_result")
    guide.paragraph("Your output shows that the pipeline completed successfully. The numbers are useful as a baseline and as proof that the scripts run, but they are not yet the project’s final 30-day findings.")
    if SCREENSHOT.exists():
        guide.figure(SCREENSHOT, "Figure 2", "The group’s terminal output after collection, normalization and analysis.", width=6.35)
    guide.table(
        ["Terminal message", "Correct interpretation"],
        [
            ("`wrote 89277 normalized observations`", "The parser produced 89,277 rows from the raw files currently in `data/raw`. These are observations/rows, not 89,277 Danish IPs and not necessarily 89,277 unique values."),
            ("`wrote 0 parser notes`", "The supported file formats were parsed without a recorded parser exception. This does not prove that an entry is malicious, Danish or currently assigned to a particular owner."),
            ("`wrote 3 daily rows`", "The analysis found three source/day summary rows: FireHOL, Spamhaus DROP and blocklist.de for the current date. This means three feed rows, not three collection days."),
            ("`wrote 44746 persistence rows`", "The persistence file contains one row for each unique source/indicator combination observed by the script. It is a working descriptive output, not a final conclusion."),
            ("`wrote 2 overlap rows`", "Only two exact representation-level source pairs were comparable in the current data: CIDR overlap between FireHOL/Spamhaus and IP overlap between FireHOL/blocklist.de."),
            ("`URLhaus is disabled`", "This is an intentional configuration message. URLhaus was not queried and is not an error in the three-source collection."),
            ("`Matplotlib is building the font cache`", "This is normal first-run plotting setup. It does not indicate a data or parser problem."),
        ],
        widths=[2.25, 4.15],
        font_size=8.9,
    )
    guide.paragraph("The analysis groups snapshots by calendar date. Because the package already contained a baseline and you collected again on the same date, the raw directory may contain more than one snapshot for that date. Keep both files for the evidence trail, but try to collect once per day at a fixed UTC time from now on. If a second run is unavoidable, document it in the collection log and do not silently treat it as a separate day.")
    guide.heading("4.1 What the current result does not tell us", 2, "current_result_limits")
    for text in [
        "It does not tell us how many indicators are Danish-related. Member 4 must add and document country evidence.",
        "It does not tell us how long an indicator remains listed. That requires observations across different days.",
        "It does not prove false positives. A plausible reassignment or country conflict is a risk signal, not proof that the listing is wrong.",
        "It does not show whether a provider will delist an entry. That requires policy analysis and, only where permitted, a controlled owned-asset case.",
        "It does not describe the whole Internet or all Danish infrastructure. The conclusion is limited to the selected feeds and observation period.",
    ]:
        guide.bullet(text)

    guide.heading("5. Your responsibility: Member 3 and Member 4", 1, "your_member_3_4_role")
    guide.paragraph("Yes. The hands-on part you want to take is mainly the combined Member 3 and Member 4 work. Member 3 preserves and processes the feed evidence. Member 4 determines which observations have defensible Danish-related evidence and adds carefully documented context. These roles depend on each other.")
    guide.heading("5.1 Member 3 — collection and normalization", 2, "member_3_tasks")
    for text in [
        "Maintain the collection schedule and run the collector once per day at the agreed UTC time.",
        "Check that each enabled source returned HTTP success, a sensible file size and a recorded SHA-256 hash.",
        "Never edit the files in `data/raw`; they are the original evidence.",
        "Run the normalizer after collection and check `logs/normalization_errors.csv`.",
        "Run the analysis script to create the current descriptive outputs.",
        "Record missing days, duplicate same-day runs, provider format changes and failed downloads.",
        "Keep the source register updated with the exact URL, policy URL, access date, terms and local filename pattern.",
    ]:
        guide.bullet(text)
    guide.heading("5.2 Member 4 — Danish evidence and enrichment", 2, "member_4_tasks")
    for text in [
        "Create a separate evidence table for allocation country, geolocation, origin ASN, reverse DNS, resolved-host country and evidence source.",
        "Use at least two independent timestamped signals for `DK-supported`.",
        "Use `DK-weak` when only one signal supports Denmark; use `Conflicting` when country signals disagree; use `Unknown` when reliable evidence is absent.",
        "Treat IP, CIDR, domain and URL attribution separately. Do not infer Danish ownership from a `.dk` suffix alone.",
        "Use passive sources first, such as RIR/RDAP, RIPEstat, DNS and permitted APIs. Query AbuseIPDB, GreyNoise or VirusTotal only for a small documented candidate set.",
        "Record the provider, retrieval time and evidence source for every enrichment result. Do not store API keys or private correspondence in the report repository.",
    ]:
        guide.bullet(text)
    guide.heading("5.3 The handoff between Members 3 and 4", 2, "member_3_4_handoff")
    guide.table(
        ["Member 3 gives Member 4", "Member 4 returns to the group"],
        [
            ("Normalized indicators with source, snapshot time, original value, hash and source path", "A classification CSV with evidence notes and DK-supported/DK-weak/Conflicting/Unknown labels"),
            ("A list of distinct IPs, CIDRs, domains and URLs requiring country review", "A documented explanation of which signal supports each classification"),
            ("Parser notes and any representation ambiguity", "A list of uncertain, conflicting or excluded cases for the limitations section"),
        ],
        widths=[3.2, 3.2],
        font_size=9.0,
    )

    guide.heading("6. The other four group roles", 1, "other_member_roles")
    guide.table(
        ["Member", "Main work", "Handoff to the whole group"],
        [
            ("Member 1", "Coordination, introduction, problem statement, aim, research questions and final coherence", "One agreed scope and one consistent explanation of the problem"),
            ("Member 2", "Literature search, literature review, theoretical framework, IEEE references and source register", "Evidence for concepts, previous findings and the research gap"),
            ("Member 3", "Collection, raw snapshots, hashes, parser, normalization and run logs", "Reproducible data pipeline and collection-quality evidence"),
            ("Member 4", "Danish classification, attribution uncertainty, passive enrichment and evidence completeness", "Country evidence and cautious interpretation"),
            ("Member 5", "Quantitative analysis, persistence, turnover, overlap and visualizations", "Tables and figures tied to the research questions"),
            ("Member 6", "Policy analysis, ethics, delisting, discussion, recommendations and conclusion", "Meaning of the findings and responsible operational use"),
        ],
        widths=[0.85, 2.7, 2.85],
        font_size=8.7,
    )
    guide.paragraph("This division is for coordination, not six isolated reports. All members should understand the full pipeline, review the complete report and take part in the final discussion and oral examination.")

    guide.heading("7. Exact hands-on road map", 1, "hands_on_road_map")
    guide.heading("7.1 Day 1 — freeze and record the baseline", 2, "day_1_baseline")
    for text in [
        "Record the date, UTC collection time, enabled sources and the exact Git/project version.",
        "Tell the group that FireHOL, Spamhaus DROP and blocklist.de were collected successfully and URLhaus was intentionally disabled.",
        "Save the terminal output or write the successful counts in the per-run collection log; the screenshot can be placed in the working notes, not used as a substitute for raw files.",
        "Confirm that `logs/normalization_errors.csv` contains only its header.",
        "Do not call the baseline a final finding. It is a smoke test and Day 1 evidence.",
    ]:
        guide.bullet(text)
    guide.heading("7.2 Daily routine", 2, "daily_routine")
    add_code(guide.doc, """cd Blocklist_Danish_Project_Package
source .venv/bin/activate

export BLOCKLIST_USER_AGENT='AAU-blocklist-study/0.1 contact: YOUR-GROUP-EMAIL'

# One command creates the raw snapshot, normalization, analysis,
# timestamped run record and independent comparison attempt.
python scripts/run_daily_pipeline.py""")
    guide.heading("7.3 Daily checks", 2, "daily_checks")
    add_code(guide.doc, """# Read the newest human-readable run handoff.
ls -td run_records/* | head -n 1
cat "$(ls -td run_records/* | head -n 1)/run_summary.md"

# Inspect timestamped events and collection status.
cat "$(ls -td run_records/* | head -n 1)/run_events.csv"
cat "$(ls -td run_records/* | head -n 1)/collection.csv"

# Confirm parser notes, dates and independent comparisons.
cat logs/normalization_errors.csv
cut -d, -f3 data/normalized/observations.csv | cut -c1-10 | sort -u
find comparisons -maxdepth 3 -type f -print | sort""")
    guide.paragraph("These commands show the dates present in the normalized file and the separate comparison artifacts. You need at least 30 distinct collection dates for the planned historical analysis. If the same date appears repeatedly, it is still one calendar day for this project’s primary comparison. If a provider is missing on one selected day, the comparison marks the source as missing and uses NA for change counts instead of calling the entries removals.")
    guide.heading("7.4 Member 4 classification step", 2, "classification_step")
    add_code(guide.doc, """cp data/danish_evidence_template.csv data/danish_evidence.csv
cp data/danish_prefixes_template.csv data/danish_prefixes.csv

# Fill both CSV files with evidence retrieved and timestamped by the group.
python scripts/classify_danish.py \\
  --input data/normalized/observations.csv \\
  --evidence data/danish_evidence.csv \\
  --prefixes data/danish_prefixes.csv \\
  --out results/danish_classification.csv""")
    guide.paragraph("The classifier does not perform network lookups. That is deliberate. Country evidence needs a source, a timestamp and a human review trail. The script turns the evidence that Member 4 has recorded into the project’s four classifications.")
    guide.heading("7.5 Safe practice work", 2, "safe_practice")
    guide.paragraph("The project does not require malicious traffic, false abuse reports or public-host scanning. If the group needs a network-security practice exercise, use only an isolated VM or a host owned and explicitly authorized by the group.")
    add_code(guide.doc, """export TARGET=192.168.56.101
ip route
ping -c 2 "$TARGET"
nmap -Pn -sT --top-ports 100 --reason "$TARGET"
nmap -Pn -sV --version-light "$TARGET""")
    guide.paragraph("Do not replace the lab target with an IP copied from a public blocklist. The project studies feed data; it does not give permission to probe listed hosts.")

    guide.heading("8. Thirty-day project plan", 1, "thirty_day_plan")
    guide.table(
        ["Time", "Member 3–4 focus", "Other group focus", "Checkpoint"],
        [
            ("Before Day 1", "Test collector, parser and evidence templates", "Agree RQs, sources, terminology and roles", "Everyone approves the method"),
            ("Days 1–3", "Collect baseline and verify hashes/status", "Capture policy pages and begin literature log", "No silent collection failures"),
            ("Days 4–7", "Run daily routine; review format consistency", "Draft Chapters 1–3", "Day 7 quality review"),
            ("Days 8–14", "Build candidate set and begin passive evidence", "Prepare analysis plan and policy coding", "Midpoint data and method review"),
            ("Days 15–21", "Continue collection; resolve uncertain classifications", "Draft findings tables and figures", "Day 21 policy review complete"),
            ("Days 22–30", "Finish coverage, hashes, classification and QA", "Draft discussion, recommendations and conclusion outline", "Freeze data after Day 30"),
            ("After Day 30", "Run final normalization and analysis", "Integrate, cross-review and prepare oral defence", "Every claim has evidence"),
        ],
        widths=[1.25, 2.2, 2.45, 1.0],
        font_size=8.45,
    )

    guide.heading("9. How to explain the project to the group", 1, "group_explanation")
    guide.paragraph("You can explain it like this:")
    add_code(guide.doc, """Our project studies public blocklists as changing security systems. We will collect
FireHOL, Spamhaus DROP and blocklist.de snapshots for at least 30 days. Every
daily run gets a UTC-timestamped manifest, command log and archived analysis.
We will preserve the original files, normalize the different formats, identify
which observations have defensible Danish evidence, and measure listing size,
additions, removals, persistence and overlap. We will compare two distinct
dates in a separate folder so a comparison error cannot change the main results.
We will then compare the operators' policies, especially correction and delisting.
We are measuring what the lists publish; we are not treating one listing as
proof of current maliciousness and we are not scanning public hosts.""")
    guide.heading("9.1 What you can say about your own role", 2, "own_role_explanation")
    guide.paragraph("I am taking the hands-on data part. I will maintain the daily collection, preserve the raw evidence, check the hashes and parser log, run normalization and analysis, and work with Member 4 on Danish classification and passive evidence. I will not make conclusions from a single day. I will give the group reproducible data and clearly label uncertainty so the final analysis is based on the complete collection period.")
    guide.heading("9.2 What the group needs to decide together", 2, "decisions_together")
    for text in [
        "The exact 30-day start and end dates and the fixed UTC collection time.",
        "Whether the three working feeds are the final minimum set and whether URLhaus is added later as a separate domain/URL source.",
        "The final Danish classification rule and the independent evidence sources accepted by the group.",
        "The definition of a successful snapshot and how failed or duplicated same-day runs will be reported.",
        "Which tables and figures answer each research question.",
        "The exact safe scope of any owned-asset delisting case, subject to provider terms and supervisor approval.",
    ]:
        guide.bullet(text)

    guide.heading("10. What remains missing", 1, "missing_parts")
    guide.paragraph("The project framework is ready, but the final empirical sections still depend on the collection period. The current baseline cannot answer the historical questions by itself.")
    guide.table(
        ["Missing item", "Who leads it", "When it becomes available"],
        [
            ("30 distinct collection dates", "Member 3", "After the daily schedule is completed"),
            ("Danish evidence and classifications", "Member 4", "After candidate indicators are reviewed"),
            ("Persistence and turnover conclusions", "Members 3 and 5", "After multiple dates, preferably all 30"),
            ("Policy/delisting comparison", "Member 6 with Member 2", "After policy pages and evidence are captured"),
            ("Final Findings chapter", "Member 5 with all members", "After QA and result review"),
            ("Final Discussion and Conclusion", "Member 6 with all members", "After the Findings chapter is frozen"),
        ],
        widths=[2.45, 1.65, 2.45],
        font_size=9.0,
    )
    guide.paragraph("Do not fill these sections with invented numbers. If a source is missing, report the missing coverage and explain its effect on the analysis.")

    guide.heading("11. Short glossary", 1, "glossary")
    guide.table(
        ["Term", "Meaning"],
        [
            ("Blocklist", "A published collection of indicators associated by an operator with a stated abuse, risk or policy category."),
            ("Feed", "The file, URL or service through which a list is distributed."),
            ("API", "A documented programmatic interface. The current collector uses public file URLs, not authenticated APIs."),
            ("Snapshot", "A copy of a source at a specific retrieval time."),
            ("Indicator", "An IP address, CIDR range, domain or URL recorded by a source."),
            ("Normalization", "Converting different source formats into a common table while preserving the original value."),
            ("Persistence", "The share of successful snapshots in which an indicator remains visible."),
            ("Turnover", "The entry and exit of indicators over time, measured through additions and removals."),
            ("Overlap", "The intersection between comparable indicator sets; Jaccard similarity is intersection divided by union."),
            ("Enrichment", "Additional context such as ASN, DNS, country, reputation or registration evidence; it is not automatic ground truth."),
            ("DK-supported", "At least two independent timestamped signals support Denmark under the group’s agreed rule."),
        ],
        widths=[1.55, 5.0],
        font_size=9.0,
    )

    guide.heading("12. Immediate next action", 1, "immediate_next_action")
    guide.paragraph("Your immediate task is not to find a dramatic result. Your task is to make the evidence collection reliable. Tell the group that the first run succeeded for three public feeds, normalization produced zero parser notes, URLhaus was intentionally disabled, and the present output is a baseline. Then agree on the fixed daily collection time and continue from Day 2.")
    guide.paragraph("After each daily run, send the group a short status message with: collection date and UTC time; source status; missing or duplicated files; parser-note count; and whether the normalized date check shows a new day. This creates a simple audit trail and makes the final methodology easy to write.")
    guide.paragraph("The final report should connect every numerical claim to a result file or figure, every policy claim to a dated source capture, and every Danish classification to recorded evidence. That is the central contribution of the hands-on work.")

    toc_entries = [(title, level, anchor) for title, level, anchor in guide.headings if title != "Table of Contents"]
    replace_placeholder_with_links(guide.doc, "[[TOC]]", toc_entries)

    output = REPORT_DIR / "Member_3_4_Hands_On_Project_Guide.docx"
    guide.doc.save(output)
    print(output)


if __name__ == "__main__":
    build_guide()
